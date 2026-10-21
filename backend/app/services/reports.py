from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.models.device import Device, InspectionTask
from app.services.device_import_columns import TEMPLATE_HEADERS
from app.services.time import utc_for_api


STATUS_FILL = {
    "normal": "C6EFCE",
    "warning": "FFEB9C",
    "critical": "F4CCCC",
    "connection_failed": "F4CCCC",
    "authentication_failed": "F4CCCC",
    "data_fetch_failed": "F4CCCC",
    "cancelled": "D9EAD3",
}
CONCLUSION_TEXT = {
    "normal": "正常",
    "warning": "警告",
    "critical": "严重",
    "connection_failed": "连接失败",
    "authentication_failed": "认证失败",
    "data_fetch_failed": "数据获取失败",
    "unsupported": "不支持",
    "not_applicable": "不适用",
    "cancelled": "已取消",
}
INSPECTION_ITEM_TEXT = {
    "device_information": "设备信息",
    "uptime": "运行时间",
    "cpu": "CPU 利用率",
    "memory": "内存利用率",
    "interfaces": "接口概览",
    "stp": "STP 状态",
    "active_alarms": "当前告警",
    "boards": "单板状态",
    "power_supplies": "电源状态",
    "fans": "风扇状态",
    "temperature": "温度",
    "interface_errors": "接口错误",
    "link_aggregation": "链路聚合",
    "routing_neighbors": "邻居状态",
    "ntp": "NTP 同步",
    "storage": "CF 卡/闪存存储",
    "configuration_saved": "配置保存状态",
}
INSPECTION_SOURCE_TEXT = {
    "unsupported": "不支持",
    "not_applicable": "不适用",
}
INSPECTION_REFERENCE_TEXT = {
    "vendor CPU metric": "厂商 CPU 指标",
    "vendor memory metric": "厂商内存指标",
    "vendor STP command or MIB": "厂商 STP 命令或 MIB",
    "vendor active alarm command or MIB": "厂商活动告警命令或 MIB",
    "vendor command or MIB": "厂商命令或 MIB",
    "ENTITY-MIB or vendor command": "ENTITY-MIB 或厂商命令",
    "ENTITY-SENSOR-MIB or vendor command": "ENTITY-SENSOR-MIB 或厂商命令",
}
INSPECTION_REASON_TEXT = {
    "No verified collector is available for this item": "该机型暂无已验证的采集方式",
    "No verified SSH command is available for this item": "该机型暂无已验证的 SSH 命令",
    "SSH output did not contain a supported value": "SSH 回显中没有可解析的数值",
    "SNMP value unavailable": "SNMP 未返回该数值，已改用 SSH 采集",
    "Device rejected the SSH command as unsupported": "设备不支持该 SSH 命令",
    "No active device alarms": "未发现设备活动告警",
    "No link aggregation is configured": "设备未配置链路聚合",
    "STP is not enabled": "设备未启用 STP",
    "OSPF is not configured or has no neighbors": "设备未启用动态路由协议或没有邻居",
    "Startup configuration file presence does not prove the current configuration is saved":
        "仅检测到启动配置文件，无法确认当前配置是否已保存",
    "SNMP sysName/sysDescr returned an invalid value": "SNMP 返回的设备名称或描述无效",
    "SNMP sysUpTime returned an invalid value": "SNMP 返回的运行时间无效",
    "SNMP interface status could not be collected": "无法采集 SNMP 接口状态",
    "SNMP value is invalid": "SNMP 返回值无效",
}

FAILURE_STAGE_TEXT = {
    "connection": "连接",
    "authentication": "认证",
    "credentials": "凭据未配置",
    "collection": "数据采集",
    "parsing": "回显解析",
    "configuration": "设备配置",
    "validation": "配置校验",
    "output_truncated": "回显超出上限",
    "host_fingerprint_untrusted": "SSH 主机密钥未确认",
    "host_fingerprint_changed": "SSH 主机密钥已变化",
    "device_deleted": "设备已删除",
    "cancelled": "任务已取消",
    "collection_error": "采集程序异常",
    "task_execution": "任务执行异常",
}


def _failure_stage_text(stage) -> str | None:
    if not stage:
        return None
    return FAILURE_STAGE_TEXT.get(str(stage), str(stage))


def _local_time(value):
    """Stored timestamps are naive UTC; the workbook shows the operator's local time."""
    if value is None:
        return None
    return utc_for_api(value).astimezone().replace(tzinfo=None)


def _header(sheet, values: list[str]) -> None:
    sheet.append(values)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")


def _excel_safe_value(value):
    """Keep external text from becoming a formula when the workbook opens."""
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def _append_row(sheet, values) -> None:
    sheet.append([_excel_safe_value(value) for value in values])


def _result_text(conclusion: str) -> str:
    return CONCLUSION_TEXT.get(conclusion, conclusion)


def _inspection_source_text(source) -> str | None:
    if source is None:
        return None
    return INSPECTION_SOURCE_TEXT.get(str(source), str(source))


def _inspection_reference_text(reference) -> str | None:
    if reference is None:
        return None
    return INSPECTION_REFERENCE_TEXT.get(str(reference), str(reference))


def _inspection_reason_text(reason) -> str | None:
    if reason is None:
        return None
    text = str(reason)
    if text.startswith("SNMP: ") and "; SSH: " in text:
        snmp_reason, ssh_reason = text.removeprefix("SNMP: ").split("; SSH: ", 1)
        return f"SNMP：{_inspection_reason_text(snmp_reason)}；SSH：{_inspection_reason_text(ssh_reason)}"
    exact = INSPECTION_REASON_TEXT.get(text)
    if exact:
        return exact
    if text.startswith("SSH command failed:"):
        return "SSH 命令执行失败：" + text.removeprefix("SSH command failed:").strip()
    if text.startswith("SSH command timed out before the device prompt was received"):
        return "SSH 命令等待设备提示符超时，后续命令未执行"
    if text.startswith("IF-MIB error counter columns unavailable:"):
        return "IF-MIB 错误计数列不可用：" + text.removeprefix(
            "IF-MIB error counter columns unavailable:"
        ).strip()
    for prefix, label in (
        ("SNMP device collection exceeded elapsed time limit", "SNMP 单设备采集超过总时间上限"),
        ("SNMP table exceeded elapsed time limit", "SNMP 表遍历超过时间上限"),
        ("SNMP table exceeded row limit", "SNMP 表遍历超过行数上限"),
    ):
        if text.startswith(prefix):
            detail = text.removeprefix(prefix).strip().strip("()")
            return label + (f"（{detail}）" if detail else "")
    if text == "SNMP table returned a non-increasing OID":
        return "SNMP 表返回了非递增 OID"
    return text


def _inspection_value_text(item_id: str, value) -> str | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return f"{value}%" if item_id in {"cpu", "memory"} else str(value)
    if isinstance(value, str):
        return value
    if not isinstance(value, dict):
        return str(value)

    def limited_details(values, formatter, limit: int = 3) -> str:
        visible = [formatter(item) for item in values[:limit]]
        if len(values) > limit:
            visible.append(f"另有 {len(values) - limit} 项")
        return "；".join(visible)

    def alarm_text(alarm: dict) -> str:
        return " / ".join(str(alarm.get(key)) for key in ("name", "severity", "description") if alarm.get(key))

    def bytes_text(raw) -> str | None:
        if not isinstance(raw, (int, float)) or raw < 0:
            return None
        amount = float(raw)
        units = ("B", "KB", "MB", "GB", "TB")
        unit = 0
        while amount >= 1024 and unit < len(units) - 1:
            amount /= 1024
            unit += 1
        displayed = str(round(amount)) if amount >= 10 or amount.is_integer() else f"{amount:.1f}"
        return f"{displayed} {units[unit]}"

    if item_id in {"cpu", "memory"} and "percent" in value:
        percent = value.get("percent")
        prefix = f"{percent}%" if percent is not None else "-"
        readings = value.get("readings") or []
        if not readings:
            return prefix

        def usage_text(reading: dict) -> str:
            object_id = "整机" if reading.get("object_id") == "system" else reading.get("object_id")
            scope = {
                "overall": "整体", "management-plane": "管理平面",
                "data-plane": "数据平面", "container": "容器",
            }.get(reading.get("scope"), reading.get("scope"))
            label = "/".join(filter(None, (
                object_id,
                scope if reading.get("scope") != "overall" else None,
            ))) or "-"
            window_text = {
                "current": "当前", "5s": "5 秒", "1m": "1 分钟", "5m": "5 分钟",
            }.get(reading.get("window"), reading.get("window"))
            window = f"（{window_text}）" if window_text else ""
            return f"{label} {reading.get('percent', '-')}%{window}"

        window_order = ("5s", "current", "1m", "5m")
        has_overall = any(reading.get("scope") == "overall" for reading in readings)
        candidates = (
            [reading for reading in readings if reading.get("scope") == "overall"]
            if has_overall else readings
        )
        selected_window = next(
            (window for window in window_order if any(
                reading.get("window") == window for reading in candidates
            )),
            candidates[0].get("window"),
        )

        def resource_priority(reading: dict) -> tuple[int, int | float]:
            candidate_scope = not has_overall or reading.get("scope") == "overall"
            if candidate_scope and reading.get("window") == selected_window:
                group = 0
            elif reading.get("window") == selected_window:
                group = 1
            else:
                try:
                    window_rank = window_order.index(reading.get("window"))
                except ValueError:
                    window_rank = len(window_order)
                group = 2 + window_rank
            percent = reading.get("percent")
            return group, -(percent if isinstance(percent, (int, float)) else -1)

        prioritized = sorted(readings, key=resource_priority)
        return f"{prefix}；{limited_details(prioritized, usage_text)}"

    if item_id == "active_alarms":
        prefix = f"共 {value.get('total', 0)} 条，严重 {value.get('critical', 0)} 条，警告 {value.get('warning', 0)} 条"
        alarms = value.get("alarms") or []
        return f"{prefix}；{limited_details(alarms, alarm_text)}" if alarms else prefix
    if item_id == "stp" and value.get("stp_enabled") is False:
        return "未启用 STP"
    if item_id == "stp":
        return (
            f"阻塞/异常端口 {value.get('stp_discarding_ports', '-')} 个，"
            f"故障端口 {value.get('stp_broken_ports', 0)} 个"
        )
    if item_id in {"boards", "power_supplies", "fans"}:
        prefix = f"共 {value.get('total', 0)} 个，异常 {value.get('abnormal', 0)} 个"
        components = value.get("components") or []
        details = []
        if components:
            def component_text(component: dict) -> str:
                detail = " / ".join(
                    str(component.get(key)) for key in ("id", "model", "status")
                    if component.get(key) and component.get(key) != "-"
                )
                alarms = component.get("alarms") or []
                return f"{detail} / {limited_details(alarms, alarm_text, 1)}" if alarms else detail

            details.append(limited_details(components, component_text))
        alarms = value.get("alarms") or []
        if alarms:
            details.append("告警 " + limited_details(alarms, alarm_text, 2))
        return f"{prefix}；{'；'.join(details)}" if details else prefix
    if item_id == "device_information":
        identity = " / ".join(
            str(value[key]) for key in ("hostname", "model", "software_version")
            if value.get(key)
        ) or None
        compatibility = value.get("adapter_compatibility")
        if not isinstance(compatibility, dict):
            return identity
        if compatibility.get("status") == "verified":
            evidence = "，".join(
                str(compatibility[key])
                for key in ("adapter_version", "verification_date")
                if compatibility.get(key)
            )
            return f"{identity or '-'}；适配已验证" + (f"（{evidence}）" if evidence else "")
        return f"{identity or '-'}；型号/版本未验证"
    if item_id == "interfaces":
        parts = [
            f"总数 {value.get('interfaces_total', '-')} 个",
            f"Up {value.get('interfaces_up', '-')} 个",
            f"Down {value.get('interfaces_down', '-')} 个",
        ]
        if value.get("interfaces_admin_down") is not None:
            parts.append(f"管理关闭 {value['interfaces_admin_down']} 个")
        if value.get("interfaces_error_disabled") is not None:
            parts.append(f"错误禁用 {value['interfaces_error_disabled']} 个")
        if value.get("interfaces_unknown") is not None:
            parts.append(f"未知 {value['interfaces_unknown']} 个")
        summary = "，".join(parts)
        details = value.get("interface_details") or []

        def interface_state(interface: dict) -> str:
            if interface.get("error_disabled"):
                return "错误禁用"
            if interface.get("admin_down") or interface.get("admin_status") == "down":
                return "管理关闭"
            state = interface.get("physical_status") or interface.get("oper_status") or "unknown"
            return {
                "up": "正常", "down": "Down", "lowerLayerDown": "下层链路 Down",
                "testing": "测试中", "dormant": "休眠", "notPresent": "未安装",
                "unknown": "状态未知",
            }.get(str(state), str(state))

        exceptional = [item for item in details if interface_state(item) != "正常"]
        if exceptional:
            summary += "；" + limited_details(
                exceptional,
                lambda item: f"{item.get('name', '-')} {interface_state(item)}",
            )
        return summary
    if item_id == "interface_errors" and "total" in value:
        prefix = f"总错误 {value.get('total', '-')}"
        if value.get("discarded") is not None:
            prefix += f"，丢弃 {value['discarded']}"
        interfaces = value.get("interfaces") or []
        if not interfaces:
            return prefix

        def interface_error_text(interface: dict) -> str:
            details = [f"总错误 {interface.get('total_errors', '-')}"]
            if interface.get("crc_errors") is not None:
                details.append(f"CRC {interface['crc_errors']}")
            if (
                interface.get("input_discards") is not None
                or interface.get("output_discards") is not None
                or interface.get("discards") is not None
            ):
                details.append(
                    "丢弃 "
                    + str(
                        (interface.get("input_discards") or 0)
                        + (interface.get("output_discards") or 0)
                        + (interface.get("discards") or 0)
                    )
                )
            return f"{interface.get('name', '-')} " + " ".join(details)

        prioritized = sorted(
            interfaces,
            key=lambda interface: sum(
                interface.get(key) if isinstance(interface.get(key), (int, float)) else 0
                for key in ("total_errors", "input_discards", "output_discards", "discards")
            ),
            reverse=True,
        )
        return f"{prefix}；{limited_details(prioritized, interface_error_text)}"
    if item_id == "temperature":
        maximum = value.get("max_current")
        prefix = f"共 {value.get('sensors', 0)} 个传感器" + (f"，最高 {maximum}°C" if maximum is not None else "")
        readings = value.get("readings") or []
        if not readings:
            return prefix

        def reading_text(reading: dict) -> str:
            thresholds = reading.get("thresholds") or {}
            threshold = "/".join(f"{key} {item}°C" for key, item in thresholds.items())
            text = f"{reading.get('name', '-')} {reading.get('current', '-')}°C {reading.get('status', 'unknown')}"
            return f"{text} 阈值 {threshold}" if threshold else text

        return f"{prefix}；{limited_details(readings, reading_text)}"
    if item_id == "storage":
        filesystems = value.get("filesystems") or []
        if filesystems:
            def filesystem_text(filesystem: dict) -> str:
                parts = [
                    f"{filesystem.get('name', '-')} 已用 {filesystem.get('used_percent', '-')}%",
                    f"总量 {bytes_text(filesystem.get('total_bytes')) or '-'}",
                    f"空闲 {bytes_text(filesystem.get('free_bytes')) or '-'}",
                ]
                if filesystem.get("access") in {"rw", "ro"}:
                    parts.append(str(filesystem["access"]))
                if filesystem.get("status") not in {None, "", "unknown"}:
                    parts.append(str(filesystem["status"]))
                return " ".join(parts)

            return "；".join(filesystem_text(filesystem) for filesystem in filesystems)
        used = value.get("used_percent")
        return f"已用 {used}%" if used is not None else None
    if item_id == "link_aggregation" and "groups" in value:
        prefix = f"共 {value.get('groups', 0)} 组，异常 {value.get('degraded', 0)} 组"
        degraded = [group for group in value.get("details") or [] if group.get("degraded")]

        def group_text(group: dict) -> str:
            members = group.get("down_members") or []
            status = f"状态 {group['status']}" if group.get("status") else "状态异常"
            return f"{group.get('name', '-')} {status}" + (f"，异常成员 {', '.join(members)}" if members else "")

        return f"{prefix}；{limited_details(degraded, group_text)}" if degraded else prefix
    if item_id == "routing_neighbors" and "total" in value:
        return f"共 {value.get('total', 0)} 个邻居，已建立 {value.get('established', 0)} 个"
    if item_id == "ntp" and value.get("configured") is False:
        return "未配置"
    if item_id == "ntp" and "synchronized" in value:
        return "已同步" if value["synchronized"] else "未同步"
    if item_id == "configuration_saved" and "saved" in value:
        return "已保存" if value["saved"] else "存在未保存配置"
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _set_widths(workbook: Workbook) -> None:
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                max(len(str(cell.value or "")) for cell in column) + 2,
                42,
            )


def create_inspection_report(db: Session, task: InspectionTask) -> BytesIO:
    del db
    workbook = Workbook()
    summary = workbook.active
    summary.title = "任务汇总"
    _header(summary, [
        "任务名称", "开始时间", "设备总数", "正常", "警告", "严重",
        "连接失败", "认证失败", "数据获取失败", "已取消", "不支持", "不适用", "失败",
    ])
    results = list(task.results)
    counts = {conclusion: sum(result.conclusion == conclusion for result in results) for conclusion in CONCLUSION_TEXT}
    _append_row(summary, [
        task.name,
        _local_time(task.started_at or task.created_at),
        task.total_devices,
        counts["normal"],
        counts["warning"],
        counts["critical"],
        counts["connection_failed"],
        counts["authentication_failed"],
        counts["data_fetch_failed"],
        counts["cancelled"],
        counts["unsupported"],
        counts["not_applicable"],
        counts["connection_failed"] + counts["authentication_failed"] + counts["data_fetch_failed"],
    ])

    devices = workbook.create_sheet("设备结果")
    _header(devices, [
        "设备名称",
        "管理IP",
        "厂商",
        "数据来源",
        "运行时间",
        "接口数量",
        "CPU(%)",
        "内存(%)",
        "结论",
        "失败阶段",
        "说明",
    ])
    for result in results:
        metrics = result.metrics or {}
        _append_row(devices, [
            result.device_name,
            result.device_ip,
            result.vendor,
            result.collection_source,
            metrics.get("uptime"),
            metrics.get("interfaces_total"),
            metrics.get("cpu"),
            metrics.get("memory"),
            _result_text(result.conclusion),
            _failure_stage_text(result.failure_stage),
            result.message,
        ])
        devices.cell(devices.max_row, 9).fill = PatternFill(
            "solid", fgColor=STATUS_FILL.get(result.conclusion, "D9EAD3")
        )

    items_sheet = workbook.create_sheet("检查项明细")
    _header(items_sheet, [
        "设备名称", "管理IP", "检查项", "检查值", "结论", "数据来源", "适配器版本", "命令或OID", "降级原因",
    ])
    for result in results:
        inspection_items = (result.metrics or {}).get("inspection_items", {})
        if not isinstance(inspection_items, dict):
            continue
        for item_id, item in inspection_items.items():
            if not isinstance(item, dict):
                continue
            _append_row(items_sheet, [
                result.device_name,
                result.device_ip,
                INSPECTION_ITEM_TEXT.get(item_id, item_id),
                _inspection_value_text(item_id, item.get("value")),
                _result_text(str(item.get("conclusion", ""))),
                _inspection_source_text(item.get("source")),
                item.get("adapter_version"),
                _inspection_reference_text(item.get("reference")),
                _inspection_reason_text(item.get("downgrade_reason")),
            ])

    exceptions = workbook.create_sheet("异常清单")
    _header(exceptions, ["管理IP", "设备名称", "结论", "失败阶段", "错误原因"])
    for result in results:
        if result.conclusion not in {"normal", "unsupported", "not_applicable"}:
            _append_row(exceptions, [
                result.device_ip,
                result.device_name,
                _result_text(result.conclusion),
                _failure_stage_text(result.failure_stage),
                result.message,
            ])

    failures = workbook.create_sheet("执行失败")
    _header(failures, ["管理IP", "设备名称", "数据来源", "失败阶段", "错误原因"])
    for result in results:
        if result.failure_stage or result.conclusion in {
            "connection_failed",
            "authentication_failed",
            "data_fetch_failed",
            "cancelled",
        }:
            _append_row(failures, [
                result.device_ip,
                result.device_name,
                result.collection_source,
                _failure_stage_text(result.failure_stage),
                result.message,
            ])

    _set_widths(workbook)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def create_device_template() -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "设备导入模板"
    _header(sheet, TEMPLATE_HEADERS)
    # The sample row uses the same wording the device form shows, so the
    # operator can copy what they see instead of guessing an English keyword.
    _append_row(sheet, [
        "核心交换机-01",
        "192.168.1.1",
        "华为",
        "huawei-vrp",
        "交换机",
        "核心",
        "",
        22,
        "ops",
        "",
        "",
        "v2c",
        161,
        "public",
        "",
        "",
        "",
        "",
        "",
        "",
    ])
    _set_widths(workbook)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def create_device_export(devices: list[Device]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "设备清单"
    _header(sheet, [
        "name",
        "ip",
        "vendor",
        "platform_id",
        "device_type",
        "group_name",
        "note",
        "ssh_port",
        "snmp_version",
        "snmp_port",
        "snmp_v3_username",
        "snmp_v3_security_level",
        "has_ssh",
        "has_snmp_credentials",
        "is_active",
    ])
    for device in devices:
        _append_row(sheet, [
            device.name,
            device.ip,
            device.vendor,
            device.platform_id,
            device.device_type,
            device.group_name,
            device.note,
            device.ssh_port,
            device.snmp_version,
            device.snmp_port,
            device.snmp_v3_username,
            device.snmp_v3_security_level,
            bool(device.ssh_username and device.ssh_password),
            bool(device.snmp_community or device.snmp_v3_username),
            device.is_active,
        ])
    _set_widths(workbook)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
