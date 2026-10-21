from io import BytesIO
from zipfile import BadZipFile, ZipFile

from datetime import datetime
import errno
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db import get_db
from app.models.device import ActivityLog, Device, InspectionTask
from app.schemas.device import DeviceCreate
from app.services.output_files import ensure_available_space, report_directory, reserved_unique_file_path, safe_filename_part
from app.services.device_import_columns import normalize_import_headers
from app.services.reports import create_device_export, create_device_template, create_inspection_report
from app.services.security import encrypt_secret
from app.services.settings import get_settings


router = APIRouter()
SECRET_FIELDS = {"ssh_password", "enable_password", "snmp_community", "snmp_v3_auth_key", "snmp_v3_priv_key"}
TERMINAL_TASK_STATUSES = {"completed", "partial_failed", "failed", "cancelled"}
MAX_EXCEL_IMPORT_SIZE = 10 * 1024 * 1024
MAX_XLSX_ENTRIES = 2000
MAX_XLSX_UNCOMPRESSED_SIZE = 100 * 1024 * 1024
MAX_XLSX_MEMBER_SIZE = 25 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 1000
MAX_WORKSHEET_ROWS = 10000
MAX_WORKSHEET_COLUMNS = 100


def _validate_xlsx_container(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ENTRIES:
                raise HTTPException(status_code=413, detail="Excel file contains too many archive entries")
            uncompressed = 0
            for entry in entries:
                size = entry.file_size
                uncompressed += size
                if size > MAX_XLSX_MEMBER_SIZE or uncompressed > MAX_XLSX_UNCOMPRESSED_SIZE:
                    raise HTTPException(status_code=413, detail="Excel file contains excessive uncompressed data")
                if entry.compress_size and size / entry.compress_size > MAX_XLSX_COMPRESSION_RATIO:
                    raise HTTPException(status_code=413, detail="Excel file compression ratio is unsafe")
    except HTTPException:
        raise
    except (BadZipFile, OSError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="Invalid Excel file container") from exc


@router.get("/template")
def download_template():
    return StreamingResponse(
        create_device_template(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=device-import-template.xlsx"},
    )


@router.get("/devices-export")
def export_devices(db: Session = Depends(get_db)):
    devices = db.query(Device).order_by(Device.ip).all()
    db.add(ActivityLog(event="device_export", message=f"Exported {len(devices)} devices without credentials"))
    db.commit()
    return StreamingResponse(
        create_device_export(devices),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=device-list.xlsx"},
    )


@router.post("/import")
async def import_devices(file: UploadFile = File(...), db: Session = Depends(get_db)):
    try:
        if not file.filename or not file.filename.lower().endswith(".xlsx"):
            raise HTTPException(status_code=400, detail="Please upload an .xlsx file")
        content = await file.read(MAX_EXCEL_IMPORT_SIZE + 1)
        if len(content) > MAX_EXCEL_IMPORT_SIZE:
            raise HTTPException(status_code=413, detail="Excel import file is too large")
        try:
            _validate_xlsx_container(content)
            workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
            sheet = workbook.active
            if sheet.max_row and sheet.max_row > MAX_WORKSHEET_ROWS:
                raise HTTPException(status_code=413, detail="Excel worksheet contains too many rows")
            if sheet.max_column and sheet.max_column > MAX_WORKSHEET_COLUMNS:
                raise HTTPException(status_code=413, detail="Excel worksheet contains too many columns")
            rows = list(sheet.iter_rows(values_only=True))
        except Exception as exc:
            if isinstance(exc, HTTPException):
                raise
            raise HTTPException(status_code=400, detail=f"Invalid Excel file: {exc}") from exc
    finally:
        if "workbook" in locals():
            workbook.close()
        await file.close()
    # Accept both the Chinese labels the template ships and the English field
    # names used by sheets exported from earlier versions.
    try:
        headers = normalize_import_headers(rows[0] if rows else ())
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    required_headers = {"ip", "vendor", "device_type"}
    missing_headers = sorted(required_headers - set(headers))
    if missing_headers:
        raise HTTPException(
            status_code=400,
            detail=f"Missing required Excel columns: {', '.join(missing_headers)}",
        )

    errors: list[dict] = []
    prepared_devices: list[tuple[int, DeviceCreate]] = []
    seen_ips: set[str] = set()
    existing_ips = {ip for (ip,) in db.query(Device.ip).all()}
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        payload = {headers[index]: value for index, value in enumerate(row) if index < len(headers) and value not in (None, "")}
        try:
            validated = DeviceCreate.model_validate(payload)
            if validated.ip in seen_ips or validated.ip in existing_ips:
                errors.append({"row": row_number, "reason": "Management IP already exists"})
            else:
                seen_ips.add(validated.ip)
                prepared_devices.append((row_number, validated))
        except Exception as exc:
            errors.append({"row": row_number, "reason": str(exc)})

    if errors:
        db.add(ActivityLog(event="device_import_failed", message=f"Rejected import with {len(errors)} invalid rows"))
        db.commit()
        return {"imported": 0, "errors": errors}

    try:
        for _, validated in prepared_devices:
            values = validated.model_dump()
            for secret_field in SECRET_FIELDS:
                if values.get(secret_field):
                    values[secret_field] = encrypt_secret(values[secret_field])
            db.add(Device(**values))
        db.add(ActivityLog(event="device_import", message=f"Imported {len(prepared_devices)} devices"))
        db.commit()
    except IntegrityError:
        db.rollback()
        return {
            "imported": 0,
            "errors": [{"row": row_number, "reason": "Management IP already exists"} for row_number, _ in prepared_devices],
        }
    return {"imported": len(prepared_devices), "errors": []}


@router.get("/report/{task_id}")
def export_report(task_id: int, db: Session = Depends(get_db)):
    task = db.get(InspectionTask, task_id)
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")
    if task.status not in TERMINAL_TASK_STATUSES or task.finished_at is None:
        raise HTTPException(status_code=409, detail="Inspection report is only available after task completion")
    # Name the file in local time so it matches the times the operator sees.
    filename = f"{safe_filename_part(task.name)}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    temporary_path = None
    try:
        content = create_inspection_report(db, task).getvalue()
        directory = report_directory(get_settings(db))
        ensure_available_space(directory, len(content))
        with reserved_unique_file_path(directory, filename) as path:
            temporary_path = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
            temporary_path.write_bytes(content)
            temporary_path.replace(path)
    except OSError as exc:
        if temporary_path:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass
        message = "磁盘可用空间不足，无法保存报告" if exc.errno == errno.ENOSPC else "Could not save inspection report"
        db.add(ActivityLog(
            level="ERROR",
            event="inspection_report_export_failed",
            task_id=task.id,
            message=f"{message}: {exc}",
        ))
        db.commit()
        raise HTTPException(status_code=507 if exc.errno == errno.ENOSPC else 500, detail=message) from exc
    db.add(ActivityLog(
        event="inspection_report_exported",
        task_id=task.id,
        message=f"Inspection report saved: {path}",
    ))
    db.commit()
    return FileResponse(
        path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=path.name,
    )

