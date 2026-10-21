from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError

from app.routers import excel
from app.schemas.device import DeviceCreate
from app.services.inspection import _crashed_device_result
from app.services.reports import create_device_export
from app.services.security import redact_configuration_text, redact_runtime_text


@pytest.mark.parametrize(
    "line, secret",
    [
        ("local-user admin password irreversible-cipher key 7 CipherText", "CipherText"),
        ("snmp-agent usm-user v3 ops authentication-mode sha key-string AuthSecret", "AuthSecret"),
        ("radius-server shared-key RadiusSecret", "RadiusSecret"),
        ("snmp-server host 10.0.0.5 traps version 2c TrapCommunity", "TrapCommunity"),
        ("crypto isakmp key VpnSecret address 10.0.0.6", "VpnSecret"),
        ("ike peer branch pre-shared-key cipher IkeSecret", "IkeSecret"),
        ("wpa2 psk pass-phrase cipher WifiSecret", "WifiSecret"),
    ],
)
def test_device_configuration_redacts_real_vendor_secret_syntax(line, secret):
    assert secret not in redact_configuration_text(line)
    assert secret not in redact_runtime_text(line)
    assert "[REDACTED]" in redact_configuration_text(line)


@pytest.mark.parametrize(
    "line, secret",
    [
        ("tacacs-server key 7 CiscoTacacsCipher", "CiscoTacacsCipher"),
        ("tacacs-server key CiscoTacacsPlain", "CiscoTacacsPlain"),
        ("radius-server key 0 CiscoRadiusPlain", "CiscoRadiusPlain"),
        ("radius-server key CiscoRadiusImplicit", "CiscoRadiusImplicit"),
        ("ip ospf message-digest-key 1 md5 7 OspfCipher", "OspfCipher"),
        ("ip ospf message-digest-key 9 md5 OspfPlain", "OspfPlain"),
    ],
)
def test_cisco_inline_keys_are_redacted_in_configuration_and_runtime(line, secret):
    for redact in (redact_configuration_text, redact_runtime_text):
        result = redact(line)
        assert secret not in result
        assert result.endswith("[REDACTED]")


@pytest.mark.parametrize(
    "line, expected",
    [
        ("tacacs-server key 7 secret", "tacacs-server key 7 [REDACTED]"),
        ("radius-server key 7 secret", "radius-server key 7 [REDACTED]"),
        (
            "ip ospf message-digest-key 1 md5 7 secret",
            "ip ospf message-digest-key 1 md5 7 [REDACTED]",
        ),
    ],
)
def test_cisco_inline_secret_placeholder_is_replaced_exactly_once(line, expected):
    assert redact_configuration_text(line) == expected
    assert redact_runtime_text(line) == expected


@pytest.mark.parametrize(
    "header, key_line, secret",
    [
        ("tacacs server TAC1", " key PlaintextSecret", "PlaintextSecret"),
        ("radius server RAD1", " key 0 RadiusPlain", "RadiusPlain"),
        ("radius-server template RAD2", " key 7 RadiusCipher", "RadiusCipher"),
        ("hwtacacs-server template TAC2", " key TacacsValue", "TacacsValue"),
    ],
)
def test_aaa_block_bare_keys_are_redacted(header, key_line, secret):
    configuration = f"{header}\n address ipv4 192.0.2.10\n{key_line}\n!\n"

    for redact in (redact_configuration_text, redact_runtime_text):
        result = redact(configuration)
        assert secret not in result
        assert key_line.rsplit(" ", 1)[0] in result
        assert "[REDACTED]" in result


def test_key_words_outside_aaa_blocks_are_not_mistaken_for_secrets():
    configuration = (
        "crypto key generate rsa modulus 2048\n"
        "key chain ROUTE_AUTH\n"
        " description ordinary line containing key material inventory\n"
    )

    assert redact_configuration_text(configuration) == configuration
    assert redact_runtime_text(configuration) == configuration


def test_unexpected_collection_errors_are_redacted_before_persistence():
    result = _crashed_device_result(
        {"id": 1, "ip": "198.51.100.1", "name": "switch", "vendor": "Huawei"},
        RuntimeError("password super-secret should not be stored"),
    )

    assert "super-secret" not in (result.message or "")
    assert "[REDACTED]" in (result.message or "")


def test_device_credentials_and_notes_have_storage_limits():
    with pytest.raises(ValidationError):
        DeviceCreate(
            name="bounded-device",
            ip="198.51.100.2",
            vendor="Huawei",
            device_type="switch",
            snmp_version="v2c",
            snmp_community="readonly",
            ssh_password="x" * 1025,
        )

    with pytest.raises(ValidationError):
        DeviceCreate(
            name="bounded-device",
            ip="198.51.100.3",
            vendor="Huawei",
            device_type="switch",
            snmp_version="v2c",
            snmp_community="readonly",
            note="x" * 2001,
        )


def test_device_export_cannot_turn_external_text_into_an_excel_formula():
    device = SimpleNamespace(
        name="=HYPERLINK(\"https://invalid.example\")",
        ip="10.0.0.1",
        vendor="Huawei",
        platform_id="huawei-vrp",
        device_type="switch",
        group_name="+cmd|' /C calc'!A0",
        note="@SUM(1,1)",
        ssh_port=22,
        snmp_version="v2c",
        snmp_port=161,
        snmp_v3_username=None,
        snmp_v3_security_level=None,
        ssh_username=None,
        ssh_password=None,
        snmp_community=None,
        is_active=True,
    )

    workbook = load_workbook(create_device_export([device]), data_only=False)
    row = list(workbook.active.iter_rows(min_row=2, max_row=2, values_only=True))[0]

    assert row[0].startswith("'=")
    assert row[5].startswith("'+")
    assert row[6].startswith("'@")


def test_xlsx_container_rejects_excessive_archive_entries(monkeypatch):
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("a", "1")
        archive.writestr("b", "2")
    monkeypatch.setattr(excel, "MAX_XLSX_ENTRIES", 1)

    with pytest.raises(HTTPException) as error:
        excel._validate_xlsx_container(output.getvalue())

    assert error.value.status_code == 413
