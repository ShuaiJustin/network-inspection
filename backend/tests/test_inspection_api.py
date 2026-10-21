import asyncio
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from datetime import datetime, timedelta, timezone
import errno
from fastapi import HTTPException
from fastapi.testclient import TestClient
from io import BytesIO
import importlib
from pathlib import Path
from threading import Barrier, Event, Thread
import time
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
import pytest
import sqlite3
from sqlalchemy import event

from app.db import SessionLocal, engine
from app import db as db_module
from app.models.device import ActivityLog, AdminUser, AppSettings, BackupRecord, Base, Device, InspectionResult, InspectionTask
from app.routers import auth as auth_router
from app.routers import backup as backup_router
from app.routers import devices as devices_router
from app.routers import excel as excel_router
from app.routers import screenshots as screenshots_router
from app.routers import settings as settings_router
from app.routers import system as system_router
from app.routers import tasks as tasks_router
from app.schemas.device import DeviceUpdate, TaskCreate
from app.schemas.settings import SettingsUpdate
from app.services import collectors, inspection, maintenance
from app.services.collectors import ProtocolResult
from app.services.reports import create_inspection_report
from main import app


client = TestClient(app)


def setup_module():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)


def _bootstrap_and_login() -> dict[str, str]:
    client.post("/auth/bootstrap", json={"username": "admin", "password": "inspect123"})
    response = client.post("/auth/login", json={"username": "admin", "password": "inspect123"})
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def test_auth_credentials_reject_a_whitespace_only_username():
    from app.routers.auth import Credentials

    with pytest.raises(ValidationError):
        Credentials.model_validate({"username": "   ", "password": "inspect123"})


def test_task_creation_rejects_more_than_500_devices():
    with pytest.raises(ValidationError):
        TaskCreate(name="Oversized task", device_ids=list(range(1, 502)))


def test_concurrent_login_failures_trigger_account_lock(monkeypatch):
    _bootstrap_and_login()
    db = SessionLocal()
    try:
        admin = db.query(AdminUser).one()
        admin.failed_attempts = 0
        admin.locked_until = None
        db.commit()
    finally:
        db.close()

    real_verify = auth_router.verify_admin_password

    def slow_verify(password_hash, password):
        if password == "definitely-wrong-password":
            time.sleep(0.1)
        return real_verify(password_hash, password)

    monkeypatch.setattr(auth_router, "verify_admin_password", slow_verify)
    request_barrier = Barrier(5)

    def wrong_login(_index):
        request_barrier.wait(timeout=5)
        return client.post(
            "/auth/login",
            json={"username": "admin", "password": "definitely-wrong-password"},
        ).status_code

    try:
        with ThreadPoolExecutor(max_workers=5) as executor:
            statuses = list(executor.map(wrong_login, range(5)))
        assert statuses == [401] * 5
        db = SessionLocal()
        try:
            admin = db.query(AdminUser).one()
            assert admin.failed_attempts == 0
            assert admin.locked_until > datetime.utcnow()
        finally:
            db.close()
        assert client.post(
            "/auth/login",
            json={"username": "admin", "password": "inspect123"},
        ).status_code == 401
    finally:
        db = SessionLocal()
        try:
            admin = db.query(AdminUser).one()
            admin.failed_attempts = 0
            admin.locked_until = None
            db.commit()
        finally:
            db.close()


def test_task_create_rejects_a_whitespace_only_name():
    with pytest.raises(ValidationError):
        TaskCreate.model_validate({"name": "   ", "device_ids": [1]})


@pytest.mark.parametrize("field", ["vendor", "device_type", "ssh_port", "snmp_port", "is_active"])
def test_device_update_rejects_null_for_required_fields(field):
    with pytest.raises(ValidationError):
        DeviceUpdate.model_validate({field: None})


@pytest.mark.parametrize("field", [
    "max_concurrency", "default_collection_mode", "snmp_timeout_seconds",
    "snmp_retries", "connection_timeout_seconds", "command_timeout_seconds",
    "network_retry_count", "cpu_warning_threshold", "cpu_critical_threshold",
    "memory_warning_threshold", "memory_critical_threshold", "log_retention_days",
    "backup_retention_count", "backup_ignore_patterns", "storage_warning_threshold",
    "storage_critical_threshold", "raw_output_retention_days",
])
def test_settings_update_rejects_null_for_required_fields(field):
    with pytest.raises(ValidationError):
        SettingsUpdate.model_validate({field: None})


def test_deleting_a_device_detaches_historical_results_and_backups():
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "History owner",
            "ip": "10.99.88.1",
            "vendor": "Huawei",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "readonly",
        },
    )
    assert created.status_code == 201
    device_id = created.json()["id"]
    db = SessionLocal()
    try:
        task = InspectionTask(name="Historical task", status="completed", total_devices=1, completed_devices=1)
        db.add(task)
        db.flush()
        result = InspectionResult(
            task_id=task.id,
            device_id=device_id,
            device_ip="10.99.88.1",
            device_name="History owner",
            vendor="Huawei",
            connection_status="success",
            authentication_status="success",
            collection_status="success",
            parsing_status="success",
            conclusion="normal",
        )
        backup = BackupRecord(
            device_id=device_id,
            device_ip="10.99.88.1",
            status="failed",
        )
        db.add_all([result, backup])
        db.commit()
        result_id, backup_id = result.id, backup.id
    finally:
        db.close()

    assert client.delete(f"/devices/{device_id}", headers=headers).status_code == 204

    db = SessionLocal()
    try:
        assert db.get(InspectionResult, result_id).device_id is None
        assert db.get(BackupRecord, backup_id).device_id is None
    finally:
        db.close()


def test_deleting_a_device_during_collection_preserves_the_completed_result(monkeypatch):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Delete-race-SW-01",
            "ip": "10.99.92.1",
            "vendor": "Huawei",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "readonly",
        },
    )
    assert created.status_code == 201
    device_id = created.json()["id"]
    db = SessionLocal()
    try:
        device = db.get(Device, device_id)
        snapshot = {
            "id": device.id,
            "name": device.name,
            "ip": device.ip,
            "vendor": device.vendor,
            "device_type": device.device_type,
            "group_name": device.group_name,
        }
        task = InspectionTask(
            name="Delete during collection",
            status="queued",
            inspection_level="quick",
            collection_mode="snmp",
            total_devices=1,
            selected_device_ids=str(device_id),
            selected_device_snapshot=[snapshot],
        )
        db.add(task)
        db.commit()
        task_id = task.id
    finally:
        db.close()

    started = asyncio.Event()
    release = asyncio.Event()

    async def blocked_inspection(current_task, device, _settings):
        started.set()
        await release.wait()
        return InspectionResult(
            task=current_task,
            device_id=device.id,
            device_ip=device.ip,
            device_name=device.name,
            vendor=device.vendor,
            collection_source="SNMP",
            connection_status="success",
            authentication_status="success",
            collection_status="success",
            parsing_status="success",
            conclusion="normal",
            metrics={},
        )

    monkeypatch.setattr(inspection, "_inspect_device", blocked_inspection)

    async def run_race():
        worker = asyncio.create_task(inspection.execute_task(task_id, SessionLocal, max_concurrency=1))
        await asyncio.wait_for(started.wait(), timeout=1)
        deletion = asyncio.create_task(
            asyncio.to_thread(client.delete, f"/devices/{device_id}", headers=headers)
        )
        await asyncio.sleep(0.05)
        release.set()
        response = await asyncio.wait_for(deletion, timeout=2)
        await asyncio.wait_for(worker, timeout=2)
        return response

    assert asyncio.run(run_race()).status_code == 204
    db = SessionLocal()
    try:
        saved_task = db.get(InspectionTask, task_id)
        saved_result = db.query(InspectionResult).filter_by(task_id=task_id).one()
        assert saved_task.status == "completed"
        assert saved_result.conclusion == "normal"
        assert saved_result.device_ip == "10.99.92.1"
        assert saved_result.device_id is None
    finally:
        db.close()


def test_tauri_desktop_origin_can_call_local_api():
    response = client.options(
        "/auth/status",
        headers={
            "Origin": "tauri://localhost",
            "Access-Control-Request-Method": "GET",
        },
    )

    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "tauri://localhost"


def test_tauri_windows_webview_origin_can_call_local_api():
    response = client.options(
        "/auth/status",
        headers={
            "Origin": "http://tauri.localhost",
            "Access-Control-Request-Method": "GET",
        },
    )

    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "http://tauri.localhost"


def test_runtime_token_protects_local_api_when_desktop_sidecar_configures_one(monkeypatch):
    monkeypatch.setattr(app.state, "runtime_token", "desktop-run-token", raising=False)

    assert client.get("/auth/status").status_code == 401
    response = client.get("/auth/status", headers={"X-Inspection-Run-Token": "desktop-run-token"})

    assert response.status_code == 200
    assert isinstance(response.json()["initialized"], bool)


def test_verified_adapter_catalog_is_read_only_and_empty_without_real_validation_evidence():
    headers = _bootstrap_and_login()

    response = client.get("/adapters/verified", headers=headers)

    assert response.status_code == 200
    assert response.json() == []


def test_adapter_capability_catalog_requires_login_and_reports_implemented_platforms():
    assert client.get("/adapters/capabilities").status_code == 401
    headers = _bootstrap_and_login()

    response = client.get("/adapters/capabilities", headers=headers)

    assert response.status_code == 200
    platforms = {item["platform_id"]: item for item in response.json()}
    assert set(platforms) == {"huawei-vrp", "h3c-comware", "ruijie-rgos", "cisco-ios"}
    assert platforms["huawei-vrp"]["status"] == "unverified"
    assert "firewall" in platforms["huawei-vrp"]["device_types"]


def test_device_catalog_is_read_from_the_backend_registry():
    headers = _bootstrap_and_login()

    response = client.get("/catalog/devices", headers=headers)

    assert response.status_code == 200
    assert [vendor["id"] for vendor in response.json()["vendors"]] == [
        "Huawei", "H3C", "Ruijie", "Cisco",
    ]
    assert response.json()["device_types"] == [
        {"id": "switch", "label": "交换机"},
        {"id": "router", "label": "路由器"},
        {"id": "firewall", "label": "防火墙"},
    ]


def test_startup_recovery_marks_interrupted_tasks_failed_and_keeps_saved_results():
    db = SessionLocal()
    try:
        task = InspectionTask(
            name="Interrupted task",
            status="running",
            total_devices=1,
            completed_devices=1,
            success_devices=1,
            selected_device_ids="",
        )
        task.results.append(InspectionResult(
            device_id=None,
            device_ip="192.0.2.9",
            device_name="Saved result",
            vendor="Huawei",
            collection_source="SNMP",
            connection_status="success",
            authentication_status="success",
            collection_status="success",
            parsing_status="success",
            conclusion="normal",
            failure_stage=None,
            message=None,
            metrics={},
            raw_output=None,
        ))
        db.add(task)
        db.commit()
        task_id = task.id
    finally:
        db.close()

    recover_interrupted_tasks = getattr(db_module, "recover_interrupted_tasks", None)
    assert callable(recover_interrupted_tasks)
    assert recover_interrupted_tasks() == 1

    db = SessionLocal()
    try:
        recovered = db.get(InspectionTask, task_id)
        assert recovered.status == "failed"
        assert recovered.finished_at is not None
        assert len(recovered.results) == 1
        assert db.query(ActivityLog).filter_by(task_id=task_id, event="task_interrupted").count() == 1
    finally:
        db.close()


def test_device_operation_lock_serializes_two_operations_for_the_same_device():
    try:
        device_operations = importlib.import_module("app.services.device_operations")
        device_operation_lock = getattr(device_operations, "device_operation_lock", None)
    except ModuleNotFoundError:
        device_operation_lock = None
    assert callable(device_operation_lock)

    async def scenario():
        first_entered = asyncio.Event()
        second_waiting = asyncio.Event()
        release_first = asyncio.Event()
        events: list[str] = []

        async def first_operation():
            async with device_operation_lock(42):
                events.append("first-entered")
                first_entered.set()
                await release_first.wait()
                events.append("first-exited")

        async def second_operation():
            await first_entered.wait()
            second_waiting.set()
            async with device_operation_lock(42):
                events.append("second-entered")

        first = asyncio.create_task(first_operation())
        second = asyncio.create_task(second_operation())
        await second_waiting.wait()
        await asyncio.sleep(0.02)
        assert events == ["first-entered"]
        release_first.set()
        await asyncio.gather(first, second)
        return events

    assert asyncio.run(scenario()) == ["first-entered", "first-exited", "second-entered"]


def test_connection_test_waits_for_an_active_operation_on_the_same_device(monkeypatch):
    from app.services.device_operations import device_operation_lock

    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Serialized probe",
            "ip": "10.99.90.242",
            "vendor": "Huawei",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "readonly",
        },
    )
    assert created.status_code == 201
    device_id = created.json()["id"]
    lock_entered = Event()
    release_lock = Event()
    probe_started = Event()
    response_holder = {}

    async def fake_probe(*_args, **_kwargs):
        probe_started.set()
        return ProtocolResult(True, "SNMP")

    async def hold_device_lock():
        async with device_operation_lock(device_id):
            lock_entered.set()
            await asyncio.to_thread(release_lock.wait, 5)

    def send_probe_request():
        response_holder["response"] = client.post(
            f"/devices/{device_id}/connection-test",
            headers=headers,
            json={"protocol": "snmp"},
        )

    monkeypatch.setattr(devices_router, "test_snmp_connectivity", fake_probe)
    holder = Thread(target=lambda: asyncio.run(hold_device_lock()))
    requester = Thread(target=send_probe_request)
    holder.start()
    assert lock_entered.wait(5)
    requester.start()
    try:
        assert not probe_started.wait(0.25)
    finally:
        release_lock.set()
        holder.join(5)
        requester.join(5)

    assert probe_started.is_set()
    assert response_holder["response"].status_code == 200


def test_connection_test_reloads_device_after_waiting_for_its_lock(monkeypatch):
    from app.services import device_operations

    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Deleted while probe waits",
            "ip": "10.99.90.243",
            "vendor": "Huawei",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "readonly",
        },
    )
    assert created.status_code == 201
    device_id = created.json()["id"]
    lock_entered = Event()
    release_lock = Event()
    collector_called = Event()
    response_holder = {}

    async def fake_probe(*_args, **_kwargs):
        collector_called.set()
        return ProtocolResult(True, "SNMP")

    async def hold_device_lock():
        async with device_operations.device_operation_lock(device_id):
            lock_entered.set()
            await asyncio.to_thread(release_lock.wait, 5)

    def send_probe_request():
        response_holder["response"] = client.post(
            f"/devices/{device_id}/connection-test",
            headers=headers,
            json={"protocol": "snmp"},
        )

    monkeypatch.setattr(devices_router, "test_snmp_connectivity", fake_probe)
    holder = Thread(target=lambda: asyncio.run(hold_device_lock()))
    requester = Thread(target=send_probe_request)
    holder.start()
    assert lock_entered.wait(5)
    requester.start()
    try:
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            entry = device_operations._device_locks.get(device_id)
            if entry is not None and entry.users == 2:
                break
            time.sleep(0.01)
        else:
            pytest.fail("connection test did not begin waiting for the device lock")
        db = SessionLocal()
        try:
            db.delete(db.get(Device, device_id))
            db.commit()
        finally:
            db.close()
    finally:
        release_lock.set()
        holder.join(5)
        requester.join(5)

    assert response_holder["response"].status_code == 404
    assert collector_called.is_set() is False


def test_device_update_waits_for_an_active_operation_on_the_same_device():
    from app.services.device_operations import device_operation_lock

    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Update-lock-SW-01",
            "ip": "10.99.92.31",
            "vendor": "Huawei",
            "device_type": "switch",
            "ssh_username": "ops",
            "ssh_password": "update-lock-password",
        },
    )
    device_id = created.json()["id"]
    lock_entered = Event()
    release_lock = Event()
    update_finished = Event()
    response_holder = {}

    async def hold_device_lock():
        async with device_operation_lock(device_id):
            lock_entered.set()
            await asyncio.to_thread(release_lock.wait, 5)

    def send_update_request():
        response_holder["response"] = client.put(
            f"/devices/{device_id}",
            headers=headers,
            json={"ssh_port": 2222},
        )
        update_finished.set()

    holder = Thread(target=lambda: asyncio.run(hold_device_lock()))
    requester = Thread(target=send_update_request)
    holder.start()
    assert lock_entered.wait(5)
    requester.start()
    try:
        assert not update_finished.wait(0.1)
    finally:
        release_lock.set()
        holder.join(5)
        requester.join(5)

    assert response_holder["response"].status_code == 200
    assert response_holder["response"].json()["ssh_port"] == 2222


def test_ssh_host_key_trust_cannot_commit_a_fingerprint_for_an_updated_port(monkeypatch):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Fingerprint-update-race-SW-01",
            "ip": "10.99.92.32",
            "vendor": "Huawei",
            "device_type": "switch",
            "ssh_username": "ops",
            "ssh_password": "fingerprint-update-race-password",
        },
    )
    device_id = created.json()["id"]
    observation_started = Event()
    release_observation = Event()
    update_finished = Event()
    responses = {}

    observed = ProtocolResult(True, "SSH")
    observed.host_fingerprint = "SHA256:old-port-fingerprint"
    observed.host_key_algorithm = "ssh-ed25519"

    async def blocked_observe(*_args, **_kwargs):
        observation_started.set()
        await asyncio.to_thread(release_observation.wait, 5)
        return observed

    monkeypatch.setattr(devices_router, "observe_ssh_host_fingerprint", blocked_observe)

    def trust_key():
        responses["trust"] = client.post(
            f"/devices/{device_id}/ssh-host-key/trust",
            headers=headers,
            json={
                "expected_host_fingerprint": observed.host_fingerprint,
                "expected_host_key_algorithm": observed.host_key_algorithm,
            },
        )

    def update_port():
        responses["update"] = client.put(
            f"/devices/{device_id}",
            headers=headers,
            json={"ssh_port": 2222},
        )
        update_finished.set()

    trust_thread = Thread(target=trust_key)
    trust_thread.start()
    assert observation_started.wait(5)
    update_thread = Thread(target=update_port)
    update_thread.start()
    try:
        assert not update_finished.wait(0.1)
    finally:
        release_observation.set()
        trust_thread.join(5)
        update_thread.join(5)

    assert responses["trust"].status_code == 200
    assert responses["update"].status_code == 200
    db = SessionLocal()
    try:
        saved = db.get(Device, device_id)
        assert saved.ssh_port == 2222
        assert saved.ssh_host_fingerprint is None
    finally:
        db.close()


def test_creating_device_hides_all_secrets_from_api_response():
    headers = _bootstrap_and_login()
    response = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Core-SW-01",
            "ip": "10.99.0.1",
            "vendor": "Huawei",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "never-return-this",
            "ssh_username": "ops",
            "ssh_password": "never-return-this-either",
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["ip"] == "10.99.0.1"
    assert "snmp_community" not in body
    assert "ssh_password" not in body


def test_rejects_duplicate_management_ip():
    headers = _bootstrap_and_login()
    first = {
        "name": "Access-SW-01",
        "ip": "10.99.0.2",
        "vendor": "H3C",
        "device_type": "switch",
        "ssh_username": "ops",
        "ssh_password": "secret",
    }

    assert client.post("/devices", headers=headers, json=first).status_code == 201
    assert client.post("/devices", headers=headers, json=first).status_code == 409


def test_snmpv3_authentication_mode_requires_complete_security_parameters():
    headers = _bootstrap_and_login()

    response = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "SNMPv3-Auth-SW-01",
            "ip": "10.99.0.4",
            "vendor": "Cisco",
            "device_type": "switch",
            "snmp_version": "v3",
            "snmp_v3_username": "readonly",
            "snmp_v3_security_level": "authNoPriv",
            "snmp_v3_auth_protocol": "SHA",
        },
    )

    assert response.status_code == 422
    assert "authNoPriv requires authentication protocol and key" in response.text


def test_snmp_only_collection_without_snmp_credentials_does_not_report_success():
    headers = _bootstrap_and_login()
    device = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Router-01",
            "ip": "10.99.0.3",
            "vendor": "Cisco",
            "device_type": "router",
            "ssh_username": "ops",
            "ssh_password": "secret",
        },
    ).json()

    response = client.post(
        "/tasks",
        headers=headers,
        json={
            "name": "SNMP only validation",
            "device_ids": [device["id"]],
            "inspection_level": "quick",
            "collection_mode": "snmp",
        },
    )

    assert response.status_code == 202
    task_id = response.json()["id"]
    task = client.get(f"/tasks/{task_id}", headers=headers).json()
    for _ in range(20):
        if task["status"] not in {"queued", "running"}:
            break
        time.sleep(0.05)
        task = client.get(f"/tasks/{task_id}", headers=headers).json()
    for _ in range(20):
        if task["status"] not in {"queued", "running"}:
            break
        time.sleep(0.05)
        task = client.get(f"/tasks/{task_id}").json()
    assert task["status"] in {"completed", "partial_failed"}
    assert task["results"][0]["conclusion"] == "data_fetch_failed"


def test_excel_import_with_an_invalid_row_is_atomic():
    headers = _bootstrap_and_login()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name", "ip", "vendor", "device_type", "snmp_version", "snmp_community"])
    sheet.append(["Valid-but-not-saved", "10.99.20.1", "Huawei", "switch", "v2c", "readonly"])
    sheet.append(["Invalid-IP", "not-an-ip", "Cisco", "router", "v2c", "readonly"])
    payload = BytesIO()
    workbook.save(payload)

    response = client.post(
        "/excel/import",
        headers=headers,
        files={
            "file": (
                "devices.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200
    assert response.json()["imported"] == 0
    assert response.json()["errors"][0]["row"] == 3
    assert response.json()["errors"][0]["reason"]
    devices = client.get("/devices", headers=headers).json()
    assert "10.99.20.1" not in {device["ip"] for device in devices}


def test_settings_are_authenticated_validated_and_persisted(tmp_path, monkeypatch):
    headers = _bootstrap_and_login()
    cleanup_calls = []
    monkeypatch.setattr(
        settings_router,
        "run_raw_output_cleanup",
        lambda session_factory, *, wait: cleanup_calls.append((session_factory, wait)),
    )

    assert client.get("/settings").status_code == 401
    defaults = client.get("/settings", headers=headers)
    assert defaults.status_code == 200
    assert defaults.json()["max_concurrency"] == 10
    assert defaults.json()["snmp_timeout_seconds"] == 3
    assert defaults.json()["storage_warning_threshold"] == 80
    assert defaults.json()["raw_output_retention_days"] == 30

    updated = client.put(
        "/settings",
        headers=headers,
        json={
            "max_concurrency": 4,
            "snmp_timeout_seconds": 5,
            "command_timeout_seconds": 45,
            "cpu_warning_threshold": 75,
            "storage_warning_threshold": 82,
            "raw_output_retention_days": 28,
            "report_directory": str(tmp_path),
        },
    )
    assert updated.status_code == 200
    assert updated.json()["max_concurrency"] == 4
    assert updated.json()["cpu_warning_threshold"] == 75
    assert updated.json()["storage_warning_threshold"] == 82
    assert updated.json()["raw_output_retention_days"] == 28
    assert cleanup_calls == [(SessionLocal, True)]

    unrelated_update = client.put(
        "/settings",
        headers=headers,
        json={"snmp_retries": 2},
    )
    assert unrelated_update.status_code == 200
    assert cleanup_calls == [(SessionLocal, True)]

    persisted = client.get("/settings", headers=headers)
    assert persisted.json()["max_concurrency"] == 4
    assert persisted.json()["report_directory"] == str(tmp_path)
    assert persisted.json()["storage_warning_threshold"] == 82
    assert persisted.json()["raw_output_retention_days"] == 28
    assert client.put("/settings", headers=headers, json={"max_concurrency": 0}).status_code == 422


def test_settings_persist_critical_thresholds_and_reject_an_inverted_range(tmp_path):
    headers = _bootstrap_and_login()

    defaults = client.get("/settings", headers=headers)
    assert defaults.status_code == 200
    assert defaults.json()["cpu_critical_threshold"] == 95
    assert defaults.json()["memory_critical_threshold"] == 95
    assert defaults.json()["storage_critical_threshold"] == 95

    updated = client.put(
        "/settings",
        headers=headers,
        json={
            "cpu_warning_threshold": 74,
            "cpu_critical_threshold": 90,
            "memory_warning_threshold": 75,
            "memory_critical_threshold": 91,
            "storage_warning_threshold": 76,
            "storage_critical_threshold": 92,
        },
    )
    assert updated.status_code == 200
    assert updated.json()["cpu_critical_threshold"] == 90
    assert updated.json()["memory_critical_threshold"] == 91
    assert updated.json()["storage_critical_threshold"] == 92

    invalid = client.put(
        "/settings",
        headers=headers,
        json={"cpu_warning_threshold": 96, "cpu_critical_threshold": 95},
    )
    assert invalid.status_code == 422


def test_settings_reject_max_concurrency_above_documented_limit():
    headers = _bootstrap_and_login()

    response = client.put("/settings", headers=headers, json={"max_concurrency": 51})

    assert response.status_code == 422


def test_settings_reject_output_directory_that_is_not_writable_directory(tmp_path):
    headers = _bootstrap_and_login()
    file_path = tmp_path / "not-a-directory.txt"
    file_path.write_text("not a directory", encoding="utf-8")

    response = client.put("/settings", headers=headers, json={"report_directory": str(file_path)})

    assert response.status_code == 422
    assert response.json()["detail"] == "Report directory is not writable"


def test_authenticated_administrator_can_change_local_password():
    headers = _bootstrap_and_login()

    changed = client.post(
        "/auth/change-password",
        headers=headers,
        json={"current_password": "inspect123", "new_password": "new-inspect-password"},
    )

    assert changed.status_code == 204
    assert client.post("/auth/login", json={"username": "admin", "password": "inspect123"}).status_code == 401
    new_login = client.post("/auth/login", json={"username": "admin", "password": "new-inspect-password"})
    assert new_login.status_code == 200
    restored = client.post(
        "/auth/change-password",
        headers={"Authorization": f"Bearer {new_login.json()['access_token']}"},
        json={"current_password": "new-inspect-password", "new_password": "inspect123"},
    )
    assert restored.status_code == 204


def test_task_cancellation_waits_for_started_device_and_marks_queue_cancelled(monkeypatch):
    db = SessionLocal()
    try:
        devices = [
            Device(name=f"Cancel-SW-{index}", ip=f"10.99.10.{index}", vendor="Huawei", device_type="switch")
            for index in range(1, 4)
        ]
        db.add_all(devices)
        db.flush()
        snapshots = [
            {
                "id": device.id,
                "name": device.name,
                "ip": device.ip,
                "vendor": device.vendor,
                "device_type": device.device_type,
                "group_name": None,
            }
            for device in devices
        ]
        task = InspectionTask(
            name="取消排队设备",
            status="queued",
            inspection_level="quick",
            collection_mode="snmp",
            total_devices=len(devices),
            selected_device_ids=",".join(str(device.id) for device in devices),
            selected_device_snapshot=snapshots,
        )
        db.add(task)
        db.commit()
        task_id = task.id
    finally:
        db.close()

    started = asyncio.Event()
    release_started_device = asyncio.Event()
    inspected_ips: list[str] = []

    async def fake_inspect(current_task, device, _settings):
        inspected_ips.append(device.ip)
        started.set()
        await release_started_device.wait()
        return InspectionResult(
            task=current_task,
            device_id=device.id,
            device_ip=device.ip,
            device_name=device.name,
            vendor=device.vendor,
            collection_source="SNMP",
            connection_status="success",
            authentication_status="success",
            collection_status="success",
            parsing_status="success",
            conclusion="normal",
            failure_stage=None,
            message=None,
            metrics={},
            raw_output=None,
        )

    monkeypatch.setattr(inspection, "inspect_device", fake_inspect)

    async def run_and_cancel():
        worker = asyncio.create_task(inspection.execute_task(task_id, SessionLocal, max_concurrency=1))
        await asyncio.wait_for(started.wait(), timeout=1)
        control_db = SessionLocal()
        try:
            current_task = control_db.get(InspectionTask, task_id)
            current_task.status = "cancelled"
            control_db.commit()
        finally:
            control_db.close()
        release_started_device.set()
        await asyncio.wait_for(worker, timeout=1)

    asyncio.run(run_and_cancel())

    assert inspected_ips == ["10.99.10.1"]
    db = SessionLocal()
    try:
        task = db.get(InspectionTask, task_id)
        results = db.query(InspectionResult).filter_by(task_id=task_id).order_by(InspectionResult.id).all()
        assert task.status == "cancelled"
        assert task.completed_devices == 3
        assert [result.conclusion for result in results] == ["normal", "cancelled", "cancelled"]
    finally:
        db.close()


def test_excel_report_has_readable_required_sheets_and_sources():
    db = SessionLocal()
    try:
        task = InspectionTask(
            name="核心设备巡检",
            status="completed",
            total_devices=1,
            completed_devices=1,
            success_devices=1,
        )
        task.results.append(
            InspectionResult(
                device_ip="10.99.30.1",
                device_name="Core-SW-01",
                vendor="Huawei",
                collection_source="SNMP",
                connection_status="success",
                authentication_status="success",
                collection_status="success",
                parsing_status="success",
                conclusion="normal",
                metrics={
                    "uptime": "4d 2h",
                    "interfaces_total": 24,
                    "cpu": 35,
                    "memory": 42,
                    "inspection_items": {
                        "cpu": {
                            "conclusion": "normal",
                            "source": "SNMP",
                            "adapter_version": "standard-snmp-v1",
                            "reference": "vendor CPU metric",
                            "downgrade_reason": None,
                            "value": 35,
                        },
                    },
                },
            )
        )
        db.add(task)
        db.commit()

        workbook = load_workbook(create_inspection_report(db, task))

        assert workbook.sheetnames == ["任务汇总", "设备结果", "检查项明细", "异常清单", "执行失败"]
        assert workbook["任务汇总"]["A1"].value == "任务名称"
        assert workbook["设备结果"]["D2"].value == "SNMP"
        assert workbook["设备结果"]["I2"].value == "正常"
        assert [cell.value for cell in workbook["检查项明细"][1]] == [
            "设备名称", "管理IP", "检查项", "检查值", "结论", "数据来源", "适配器版本", "命令或OID", "降级原因",
        ]
        assert [cell.value for cell in workbook["检查项明细"][2]] == [
            "Core-SW-01", "10.99.30.1", "CPU 利用率", "35%", "正常", "SNMP", "standard-snmp-v1", "厂商 CPU 指标", None,
        ]
    finally:
        db.close()


def test_excel_report_separates_all_task_conclusion_counts():
    conclusions = [
        "normal",
        "warning",
        "critical",
        "connection_failed",
        "authentication_failed",
        "data_fetch_failed",
        "cancelled",
        "unsupported",
        "not_applicable",
    ]
    db = SessionLocal()
    try:
        task = InspectionTask(name="状态汇总", status="completed", total_devices=len(conclusions))
        for index, conclusion in enumerate(conclusions, start=1):
            task.results.append(InspectionResult(
                device_ip=f"192.0.2.{index}",
                device_name=f"Device-{index}",
                vendor="Cisco",
                collection_source="SNMP",
                connection_status="success",
                authentication_status="success",
                collection_status="success",
                parsing_status="success",
                conclusion=conclusion,
                metrics={},
            ))
        db.add(task)
        db.commit()

        summary = load_workbook(create_inspection_report(db, task))["任务汇总"]

        assert [cell.value for cell in summary[1]] == [
            "任务名称", "开始时间", "设备总数", "正常", "警告", "严重",
            "连接失败", "认证失败", "数据获取失败", "已取消", "不支持", "不适用", "失败",
        ]
        values = [cell.value for cell in summary[2]]
        assert values[0] == "状态汇总"
        expected_start = task.created_at.replace(tzinfo=timezone.utc).astimezone().replace(tzinfo=None)
        assert values[1].replace(microsecond=0) == expected_start.replace(microsecond=0)
        assert values[2:] == [9, 1, 1, 1, 1, 1, 1, 1, 1, 1, 3]
    finally:
        db.close()


def test_task_response_separates_device_conclusion_counts():
    task = InspectionTask(name="任务状态", total_devices=4)
    for conclusion in ["normal", "warning", "authentication_failed", "cancelled"]:
        task.results.append(InspectionResult(
            device_ip="192.0.2.1",
            device_name="Device",
            vendor="Huawei",
            collection_source="SNMP",
            connection_status="success",
            authentication_status="success",
            collection_status="success",
            parsing_status="success",
            conclusion=conclusion,
            metrics={},
        ))

    response = tasks_router._task_response(task)

    assert response["normal_devices"] == 1
    assert response["warning_devices"] == 1
    assert response["authentication_failed_devices"] == 1
    assert response["cancelled_devices"] == 1
    assert response["critical_devices"] == 0
    assert response["connection_failed_devices"] == 0
    assert response["data_fetch_failed_devices"] == 0
    assert response["unsupported_devices"] == 0
    assert response["not_applicable_devices"] == 0


def test_task_list_aggregates_counts_without_loading_result_payloads():
    db = SessionLocal()
    task = InspectionTask(name="List query audit", status="completed", total_devices=1)
    task.results.append(InspectionResult(
        device_ip="192.0.2.199",
        device_name="Query audit",
        vendor="Huawei",
        collection_source="SNMP",
        connection_status="success",
        authentication_status="success",
        collection_status="success",
        parsing_status="success",
        conclusion="normal",
        metrics={"large": "payload"},
        raw_output="must not be loaded by the list endpoint",
    ))
    db.add(task)
    db.commit()

    statements = []

    def capture_statement(_connection, _cursor, statement, _parameters, _context, _many):
        statements.append(statement.lower())

    event.listen(engine, "before_cursor_execute", capture_statement)
    try:
        tasks_router.list_tasks(db)
    finally:
        event.remove(engine, "before_cursor_execute", capture_statement)
        db.delete(task)
        db.commit()
        db.close()

    result_queries = [statement for statement in statements if "inspection_results" in statement]
    assert result_queries
    assert any("count(" in statement for statement in result_queries)
    assert all("metrics" not in statement for statement in result_queries)
    assert all("raw_output" not in statement for statement in result_queries)


def test_snmp_authentication_error_classifier_is_shared_by_connectivity_and_collection():
    assert collectors._snmp_error_is_authentication("authorizationError") is True
    assert collectors._snmp_error_is_authentication("noAccess") is True
    assert collectors._snmp_error_is_authentication("Unknown user name") is True
    assert collectors._snmp_error_is_authentication("wrongDigest") is True
    assert collectors._snmp_error_is_authentication("noSuchName") is False


def test_inspection_uses_saved_snmp_options_and_thresholds(monkeypatch):
    observed: dict[str, int] = {}
    device = Device(
        name="Core-SW-02",
        ip="10.99.40.1",
        vendor="Huawei",
        device_type="switch",
        snmp_version="v2c",
        snmp_community="encrypted-not-used-by-test",
    )
    task = InspectionTask(name="设置生效", collection_mode="snmp", inspection_level="quick")
    settings = AppSettings(
        snmp_timeout_seconds=7,
        snmp_retries=2,
        cpu_warning_threshold=75,
        memory_warning_threshold=85,
    )

    async def fake_collect_snmp(_device, timeout, retries):
        observed["timeout"] = timeout
        observed["retries"] = retries
        return ProtocolResult(
            ok=True,
            source="SNMP",
            metrics={"hostname": "Core-SW-02", "cpu": 76, "memory": 70},
        )

    monkeypatch.setattr(inspection, "collect_snmp", fake_collect_snmp)

    result = asyncio.run(inspection.inspect_device(task, device, settings))

    assert observed == {"timeout": 7, "retries": 2}
    assert result.conclusion == "warning"


def test_logs_are_authenticated_filterable_and_paginated():
    headers = _bootstrap_and_login()
    db = SessionLocal()
    try:
        db.add_all([
            ActivityLog(level="INFO", event="task_started", device_ip="10.99.50.1", message="log-filter-marker started"),
            ActivityLog(level="WARNING", event="connection_failed", device_ip="10.99.50.2", message="log-filter-marker timeout"),
        ])
        db.commit()
    finally:
        db.close()

    assert client.get("/logs").status_code == 401
    response = client.get("/logs", headers=headers, params={"level": "WARNING", "keyword": "marker", "limit": 10, "offset": 0})

    assert response.status_code == 200
    assert response.json()["total"] == 1
    assert response.json()["items"][0]["event"] == "connection_failed"
    assert response.json()["items"][0]["device_ip"] == "10.99.50.2"


def test_delete_logs_removes_only_selected_ids():
    headers = _bootstrap_and_login()
    db = SessionLocal()
    try:
        selected = ActivityLog(event="delete-selected", message="selected-log-marker")
        retained = ActivityLog(event="delete-selected", message="retained-log-marker")
        db.add_all([selected, retained])
        db.commit()
        selected_id = selected.id
        retained_id = retained.id
    finally:
        db.close()

    response = client.request("DELETE", "/logs", headers=headers, json={"ids": [selected_id]})

    assert response.status_code == 200
    assert response.json() == {"deleted": 1}
    db = SessionLocal()
    try:
        assert db.get(ActivityLog, selected_id) is None
        assert db.get(ActivityLog, retained_id) is not None
    finally:
        db.close()


def test_delete_logs_uses_filter_and_snapshot_boundary():
    headers = _bootstrap_and_login()
    db = SessionLocal()
    try:
        old = ActivityLog(level="WARNING", event="bulk-delete", message="bulk-delete-marker old")
        db.add(old)
        db.commit()
        before_id = old.id
        newer = ActivityLog(level="WARNING", event="bulk-delete", message="bulk-delete-marker new")
        different = ActivityLog(level="INFO", event="bulk-delete", message="bulk-delete-marker different")
        db.add_all([newer, different])
        db.commit()
        newer_id = newer.id
        different_id = different.id
    finally:
        db.close()

    response = client.request(
        "DELETE",
        "/logs",
        headers=headers,
        json={
            "delete_all": True,
            "before_id": before_id,
            "level": "WARNING",
            "keyword": "bulk-delete-marker",
        },
    )

    assert response.status_code == 200
    assert response.json() == {"deleted": 1}
    db = SessionLocal()
    try:
        assert db.get(ActivityLog, before_id) is None
        assert db.get(ActivityLog, newer_id) is not None
        assert db.get(ActivityLog, different_id) is not None
    finally:
        db.close()


def test_delete_logs_rejects_an_empty_scope():
    headers = _bootstrap_and_login()

    response = client.request("DELETE", "/logs", headers=headers, json={})

    assert response.status_code == 422


def test_activity_log_database_retention_removes_only_expired_rows():
    from app.services.log_maintenance import prune_activity_log_records

    now = datetime.utcnow()
    db = SessionLocal()
    try:
        expired = ActivityLog(event="retention-expired", created_at=now - timedelta(days=3))
        retained = ActivityLog(event="retention-retained", created_at=now - timedelta(hours=3))
        db.add_all([expired, retained])
        db.commit()
        expired_id = expired.id
        retained_id = retained.id

        removed = prune_activity_log_records(db, retention_days=1, now=now, batch_size=100)

        assert removed >= 1
        assert db.get(ActivityLog, expired_id) is None
        assert db.get(ActivityLog, retained_id) is not None
    finally:
        db.close()


def test_activity_logs_are_written_to_daily_redacted_files_and_pruned(monkeypatch, tmp_path):
    monkeypatch.setattr(db_module, "ACTIVITY_LOG_DIR", tmp_path, raising=False)
    expired_log = tmp_path / "inspection_2000-01-01.log"
    expired_log.write_text("expired", encoding="utf-8")

    db = SessionLocal()
    try:
        settings = db.get(AppSettings, 1)
        if settings is None:
            settings = AppSettings(id=1)
            db.add(settings)
        settings.log_retention_days = 1
        db.add(ActivityLog(
            level="ERROR",
            event="connection_failed",
            device_ip="10.99.50.3",
            task_id=101,
            message="SSH password secret-value rejected",
        ))
        db.commit()
    finally:
        db.close()

    daily_logs = list(tmp_path.glob("inspection_*.log"))
    assert len(daily_logs) == 1
    content = daily_logs[0].read_text(encoding="utf-8")
    assert '"event": "connection_failed"' in content
    assert '"device_ip": "10.99.50.3"' in content
    assert "secret-value" not in content
    assert "[REDACTED]" in content
    assert not expired_log.exists()


def test_backup_diff_redacts_credential_lines(tmp_path, monkeypatch):
    headers = _bootstrap_and_login()
    monkeypatch.setattr(
        backup_router,
        "get_settings",
        lambda _db: SimpleNamespace(backup_directory=str(tmp_path), backup_ignore_patterns=[]),
    )
    previous_path = tmp_path / "previous.txt"
    current_path = tmp_path / "current.txt"
    previous_path.write_text("sysname Core-SW\nsnmp-agent community read old-community\n", encoding="utf-8")
    current_path.write_text(
        "sysname Core-SW\nsnmp-agent community read new-community\nlocal-user ops password simple super-secret\n",
        encoding="utf-8",
    )
    db = SessionLocal()
    try:
        device = Device(name="Backup-Test", ip="10.99.60.1", vendor="Huawei", device_type="switch")
        db.add(device)
        db.flush()
        db.add(BackupRecord(device_id=device.id, device_ip=device.ip, file_path=str(previous_path), status="success"))
        current = BackupRecord(device_id=device.id, device_ip=device.ip, file_path=str(current_path), status="success")
        db.add(current)
        db.commit()
        backup_id = current.id
    finally:
        db.close()

    response = client.get(f"/backup/{backup_id}/diff", headers=headers)

    assert response.status_code == 200
    output = "\n".join(response.json()["diff"])
    assert "new-community" not in output
    assert "super-secret" not in output
    assert "[REDACTED]" in output


def test_backup_file_is_redacted_before_writing_to_disk(monkeypatch, tmp_path):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Backup-Redaction-SW-01",
            "ip": "10.99.60.2",
            "vendor": "Huawei",
            "device_type": "switch",
            "ssh_username": "ops",
            "ssh_password": "backup-password",
        },
    )
    assert created.status_code == 201
    assert client.put("/settings", headers=headers, json={"backup_directory": str(tmp_path)}).status_code == 200

    async def fake_backup(*_args, **_kwargs):
        return ProtocolResult(
            True,
            "SSH",
            raw_output="sysname Core-SW\nsnmp-agent community read private-community\nlocal-user ops password simple secret-value\nreturn\n",
        )

    monkeypatch.setattr(backup_router, "backup_configuration", fake_backup)
    response = client.post("/backup", headers=headers, json=[created.json()["id"]])

    assert response.status_code == 200
    saved_content = next(tmp_path.rglob("*.txt")).read_text(encoding="utf-8")
    assert "private-community" not in saved_content
    assert "secret-value" not in saved_content
    assert "[REDACTED]" in saved_content


def test_backup_same_second_creates_distinct_history_files(monkeypatch, tmp_path):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Backup-History-SW",
            "ip": "10.99.60.4",
            "vendor": "Huawei",
            "device_type": "switch",
            "ssh_username": "ops",
            "ssh_password": "backup-history-password",
        },
    )
    assert created.status_code == 201
    assert client.put("/settings", headers=headers, json={"backup_directory": str(tmp_path)}).status_code == 200

    class FrozenDatetime:
        @classmethod
        def now(cls):
            return datetime(2026, 8, 1, 9, 0, 0)

    outputs = iter([
        "sysname Core-SW\ninterface Vlanif1\n description revision one\nreturn\n",
        "sysname Core-SW\ninterface Vlanif1\n description revision two\nreturn\n",
    ])

    async def fake_backup(*_args, **_kwargs):
        return ProtocolResult(True, "SSH", raw_output=next(outputs))

    monkeypatch.setattr(backup_router, "datetime", FrozenDatetime)
    monkeypatch.setattr(backup_router, "backup_configuration", fake_backup)
    first = client.post("/backup", headers=headers, json=[created.json()["id"]])
    second = client.post("/backup", headers=headers, json=[created.json()["id"]])

    assert first.status_code == 200
    assert second.status_code == 200
    files = sorted(tmp_path.rglob("*.txt"))
    assert len(files) == 2
    assert {path.read_text(encoding="utf-8") for path in files} == {
        "sysname Core-SW\ninterface Vlanif1\n description revision one\nreturn\n",
        "sysname Core-SW\ninterface Vlanif1\n description revision two\nreturn\n",
    }


def test_backup_rejects_insufficient_disk_space_before_writing(monkeypatch, tmp_path):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Backup-Disk-Space-SW",
            "ip": "10.99.60.8",
            "vendor": "Huawei",
            "device_type": "switch",
            "ssh_username": "ops",
            "ssh_password": "backup-disk-space-password",
        },
    )
    assert created.status_code == 201
    assert client.put("/settings", headers=headers, json={"backup_directory": str(tmp_path)}).status_code == 200

    async def fake_backup(*_args, **_kwargs):
        return ProtocolResult(True, "SSH", raw_output="sysname Core-SW\ninterface Vlanif1\nreturn\n")

    def no_space(*_args, **_kwargs):
        raise OSError(errno.ENOSPC, "No space left on device")

    monkeypatch.setattr(backup_router, "backup_configuration", fake_backup)
    monkeypatch.setattr(backup_router, "ensure_available_space", no_space, raising=False)
    response = client.post("/backup", headers=headers, json=[created.json()["id"]])

    assert response.status_code == 200
    assert response.json()["records"] == [{
        "device_ip": "10.99.60.8",
        "status": "failed",
        "message": "磁盘可用空间不足，无法保存备份",
    }]
    assert not list(tmp_path.rglob("*.txt"))
    assert not list(tmp_path.rglob("*.tmp"))
    db = SessionLocal()
    try:
        assert db.query(ActivityLog).filter_by(event="backup_save_failed", device_ip="10.99.60.8").count() == 1
    finally:
        db.close()


def test_backup_rejects_error_output_without_creating_a_valid_record(monkeypatch, tmp_path):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Backup-Error-SW-01",
            "ip": "10.99.60.7",
            "vendor": "Huawei",
            "device_type": "switch",
            "ssh_username": "ops",
            "ssh_password": "backup-error-password",
        },
    )
    assert created.status_code == 201
    assert client.put("/settings", headers=headers, json={"backup_directory": str(tmp_path)}).status_code == 200

    async def fake_backup(*_args, **_kwargs):
        return ProtocolResult(
            True,
            "SSH",
            raw_output="Error: display current-configuration is not permitted for this account\n",
        )

    monkeypatch.setattr(backup_router, "backup_configuration", fake_backup)
    response = client.post("/backup", headers=headers, json=[created.json()["id"]])

    assert response.status_code == 200
    assert response.json()["records"] == [{
        "device_ip": "10.99.60.7",
        "status": "failed",
        "message": "Configuration output contains a device error message",
    }]
    assert not list(tmp_path.rglob("*.txt"))


def test_inspection_result_raw_output_is_redacted_before_storage(monkeypatch):
    device = Device(
        name="Raw-Output-SW-01",
        ip="10.99.60.3",
        vendor="Cisco",
        device_type="switch",
        snmp_version="v2c",
        snmp_community="encrypted-not-used-by-test",
    )
    task = InspectionTask(name="原始输出脱敏", collection_mode="snmp", inspection_level="quick")

    async def fake_collect_snmp(*_args, **_kwargs):
        return ProtocolResult(
            True,
            "SNMP",
            metrics={"hostname": "Raw-Output-SW-01"},
            raw_output="password secret-value\ncommunity private-community\n",
        )

    monkeypatch.setattr(inspection, "collect_snmp", fake_collect_snmp)
    result = asyncio.run(inspection.inspect_device(task, device, AppSettings()))

    assert "secret-value" not in (result.raw_output or "")
    assert "private-community" not in (result.raw_output or "")
    assert "[REDACTED]" in (result.raw_output or "")


def test_changed_ssh_host_fingerprint_blocks_collection_before_authentication(monkeypatch):
    from app.services import collectors as collectors_service

    device = Device(
        name="Fingerprint-SW-01",
        ip="10.99.60.4",
        vendor="Cisco",
        device_type="switch",
        ssh_username="ops",
        ssh_password="encrypted-not-used-by-test",
    )
    device.ssh_host_fingerprint = "SHA256:old-fingerprint"
    monkeypatch.setattr(
        collectors_service,
        "_read_ssh_host_fingerprint",
        lambda *_args, **_kwargs: ("ssh-ed25519", "SHA256:new-fingerprint"),
        raising=False,
    )
    monkeypatch.setattr(collectors_service, "decrypt_secret", lambda _value: "test-password")

    def should_not_connect(**_kwargs):
        raise AssertionError("SSH authentication must not begin after a fingerprint mismatch")

    monkeypatch.setattr(collectors_service, "ConnectHandler", should_not_connect)
    result = asyncio.run(collectors_service.collect_ssh(device, "quick"))

    assert result.failure_stage == "host_fingerprint_changed"
    assert result.host_fingerprint == "SHA256:new-fingerprint"
    assert result.host_key_algorithm == "ssh-ed25519"
    assert "confirmation" in (result.message or "").lower()


def test_untrusted_ssh_host_key_blocks_collection_before_authentication(monkeypatch):
    from app.services import collectors as collectors_service

    device = Device(
        name="Untrusted-Fingerprint-SW-01",
        ip="10.99.60.5",
        vendor="Cisco",
        device_type="switch",
        ssh_username="ops",
        ssh_password="encrypted-not-used-by-test",
    )
    monkeypatch.setattr(
        collectors_service,
        "_read_ssh_host_fingerprint",
        lambda *_args, **_kwargs: ("rsa-sha2-512", "SHA256:first-fingerprint"),
        raising=False,
    )
    monkeypatch.setattr(collectors_service, "decrypt_secret", lambda _value: "test-password")

    def should_not_connect(**_kwargs):
        raise AssertionError("SSH authentication must not begin before first host-key trust")

    monkeypatch.setattr(collectors_service, "ConnectHandler", should_not_connect)
    result = asyncio.run(collectors_service.collect_ssh(device, "quick"))

    assert result.failure_stage == "host_fingerprint_untrusted"
    assert result.host_key_algorithm == "rsa-sha2-512"
    assert result.host_fingerprint == "SHA256:first-fingerprint"


def test_successful_ssh_connection_test_does_not_automatically_trust_host_key(monkeypatch):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Fingerprint-Record-SW-01",
            "ip": "10.99.60.5",
            "vendor": "Cisco",
            "device_type": "switch",
            "ssh_username": "ops",
            "ssh_password": "fingerprint-password",
        },
    )
    assert created.status_code == 201
    observed = ProtocolResult(True, "SSH")
    observed.host_fingerprint = "SHA256:first-fingerprint"
    observed.host_key_algorithm = "rsa-sha2-512"

    async def fake_test_ssh(*_args, **_kwargs):
        return observed

    monkeypatch.setattr(devices_router, "test_ssh_connectivity", fake_test_ssh)
    response = client.post(
        f"/devices/{created.json()['id']}/connection-test",
        headers=headers,
        json={"protocol": "ssh"},
    )

    assert response.status_code == 200
    assert response.json()["host_key_algorithm"] == "rsa-sha2-512"
    db = SessionLocal()
    try:
        saved = db.get(Device, created.json()["id"])
        assert getattr(saved, "ssh_host_fingerprint", None) is None
    finally:
        db.close()


def test_administrator_can_confirm_changed_ssh_host_fingerprint(monkeypatch):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Fingerprint-Confirm-SW-01",
            "ip": "10.99.60.6",
            "vendor": "Cisco",
            "device_type": "switch",
            "ssh_username": "ops",
            "ssh_password": "fingerprint-confirm-password",
        },
    )
    assert created.status_code == 201
    db = SessionLocal()
    try:
        device = db.get(Device, created.json()["id"])
        device.ssh_host_fingerprint = "SHA256:old-fingerprint"
        db.commit()
    finally:
        db.close()

    observed = ProtocolResult(True, "SSH")
    observed.host_fingerprint = "SHA256:confirmed-new-fingerprint"
    observed.host_key_algorithm = "ssh-ed25519"

    async def fake_observe(*_args, **_kwargs):
        return observed

    monkeypatch.setattr(devices_router, "observe_ssh_host_fingerprint", fake_observe, raising=False)
    response = client.post(
        f"/devices/{created.json()['id']}/ssh-host-key/trust",
        headers=headers,
        json={
            "expected_host_fingerprint": "SHA256:confirmed-new-fingerprint",
            "expected_host_key_algorithm": "ssh-ed25519",
        },
    )

    assert response.status_code == 200
    db = SessionLocal()
    try:
        assert db.get(Device, created.json()["id"]).ssh_host_fingerprint == "SHA256:confirmed-new-fingerprint"
    finally:
        db.close()


def test_ssh_host_key_trust_rejects_a_key_that_changed_after_confirmation(monkeypatch):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Fingerprint-Race-SW-01",
            "ip": "10.99.60.10",
            "vendor": "Cisco",
            "device_type": "switch",
            "ssh_username": "ops",
            "ssh_password": "fingerprint-race-password",
        },
    )
    assert created.status_code == 201
    observed = ProtocolResult(True, "SSH")
    observed.host_fingerprint = "SHA256:replacement-fingerprint"
    observed.host_key_algorithm = "ssh-ed25519"

    async def fake_observe(*_args, **_kwargs):
        return observed

    monkeypatch.setattr(devices_router, "observe_ssh_host_fingerprint", fake_observe, raising=False)
    response = client.post(
        f"/devices/{created.json()['id']}/ssh-host-key/trust",
        headers=headers,
        json={
            "expected_host_fingerprint": "SHA256:shown-fingerprint",
            "expected_host_key_algorithm": "ssh-ed25519",
        },
    )

    assert response.status_code == 409
    db = SessionLocal()
    try:
        assert db.get(Device, created.json()["id"]).ssh_host_fingerprint is None
    finally:
        db.close()


def test_backup_retention_removes_old_files_and_records(tmp_path):
    from app.routers.backup import prune_expired_backups

    db = SessionLocal()
    try:
        device = Device(name="Retention-Test", ip="10.99.70.1", vendor="Huawei", device_type="switch")
        db.add(device)
        db.flush()
        records = []
        for index in range(3):
            path = tmp_path / f"backup-{index}.txt"
            path.write_text(f"configuration {index}", encoding="utf-8")
            record = BackupRecord(
                device_id=device.id,
                device_ip=device.ip,
                file_path=str(path),
                status="success",
                created_at=datetime.utcnow() - timedelta(days=3 - index),
            )
            db.add(record)
            records.append(record)
        db.commit()

        prune_expired_backups(db, device.id, keep_count=2, root_dir=tmp_path)
        db.commit()

        remaining = db.query(BackupRecord).filter(BackupRecord.device_id == device.id).order_by(BackupRecord.id).all()
        assert len(remaining) == 2
        assert not (tmp_path / "backup-0.txt").exists()
        assert (tmp_path / "backup-1.txt").exists()
        assert (tmp_path / "backup-2.txt").exists()
    finally:
        db.close()


def test_configuration_backup_rejects_an_empty_device_selection():
    response = client.post("/backup", headers=_bootstrap_and_login(), json=[])

    assert response.status_code == 422


def test_deleting_a_device_during_backup_keeps_the_committed_backup(monkeypatch, tmp_path):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Backup-delete-race-SW-01",
            "ip": "10.99.92.2",
            "vendor": "Huawei",
            "device_type": "switch",
            "ssh_username": "ops",
            "ssh_password": "backup-delete-race-password",
        },
    )
    assert created.status_code == 201
    device_id = created.json()["id"]
    assert client.put(
        "/settings", headers=headers, json={"backup_directory": str(tmp_path)}
    ).status_code == 200
    started = Event()
    release = Event()
    deletion_finished = Event()
    responses = {}
    backup_errors = []
    validation_observed_delete = []
    real_validate_backup_output = backup_router._validate_backup_output

    async def blocked_backup(_device, _connection_timeout, _command_timeout):
        started.set()
        await asyncio.to_thread(release.wait, 5)
        return ProtocolResult(
            True,
            "SSH",
            raw_output="sysname Core-SW\ninterface Vlanif1\nreturn\n",
        )

    def wait_for_concurrent_delete(output, vendor=None, platform_id=None):
        validation_observed_delete.append(deletion_finished.wait(0.5))
        return real_validate_backup_output(output, vendor, platform_id)

    monkeypatch.setattr(backup_router, "backup_configuration", blocked_backup)
    monkeypatch.setattr(backup_router, "_validate_backup_output", wait_for_concurrent_delete)

    def run_backup():
        try:
            responses["backup"] = client.post("/backup", headers=headers, json=[device_id])
        except Exception as exc:
            backup_errors.append(exc)

    def run_delete():
        responses["delete"] = client.delete(f"/devices/{device_id}", headers=headers)
        deletion_finished.set()

    backup_thread = Thread(target=run_backup)
    delete_thread = Thread(target=run_delete)
    backup_thread.start()
    assert started.wait(5)
    delete_thread.start()
    try:
        assert not deletion_finished.wait(0.1)
    finally:
        release.set()
        backup_thread.join(5)
        delete_thread.join(5)

    assert not backup_thread.is_alive()
    assert not delete_thread.is_alive()
    assert backup_errors == []
    assert validation_observed_delete == [False]
    assert responses["backup"].status_code == 200
    assert responses["delete"].status_code == 204
    db = SessionLocal()
    try:
        record = db.query(BackupRecord).filter_by(device_ip="10.99.92.2").one()
        assert record.status == "success"
        assert record.device_id is None
        assert Path(record.file_path).is_file()
    finally:
        db.close()


def test_backup_retention_never_deletes_a_file_outside_the_backup_root(tmp_path, monkeypatch):
    allowed_root = tmp_path / "allowed-backups"
    allowed_root.mkdir()
    outside_file = tmp_path / "must-not-be-deleted.txt"
    outside_file.write_text("external data", encoding="utf-8")
    inside_file = allowed_root / "current.txt"
    inside_file.write_text("current configuration", encoding="utf-8")
    monkeypatch.setattr(backup_router, "BACKUP_DIR", allowed_root)

    db = SessionLocal()
    try:
        device = Device(name="Path-Safety", ip="10.99.70.10", vendor="Huawei", device_type="switch")
        db.add(device)
        db.flush()
        db.add(BackupRecord(
            device_id=device.id,
            device_ip=device.ip,
            file_path=str(outside_file),
            status="success",
            created_at=datetime.utcnow() - timedelta(days=2),
        ))
        db.add(BackupRecord(
            device_id=device.id,
            device_ip=device.ip,
            file_path=str(inside_file),
            status="success",
            created_at=datetime.utcnow(),
        ))
        db.commit()

        backup_router.prune_expired_backups(db, device.id, keep_count=1)
        db.commit()

        assert outside_file.exists()
        assert db.query(BackupRecord).filter(BackupRecord.device_id == device.id).count() == 1
    finally:
        db.close()


def test_backup_diff_rejects_record_paths_outside_the_backup_root(tmp_path, monkeypatch):
    headers = _bootstrap_and_login()
    allowed_root = tmp_path / "allowed-backups"
    allowed_root.mkdir()
    previous_path = tmp_path / "outside-previous.txt"
    current_path = tmp_path / "outside-current.txt"
    previous_path.write_text("sysname Before\n", encoding="utf-8")
    current_path.write_text("sysname After\n", encoding="utf-8")
    monkeypatch.setattr(
        backup_router,
        "get_settings",
        lambda _db: SimpleNamespace(backup_directory=str(allowed_root), backup_ignore_patterns=[]),
    )
    db = SessionLocal()
    try:
        device = Device(name="Unsafe-Diff", ip="10.99.70.11", vendor="Huawei", device_type="switch")
        db.add(device)
        db.flush()
        db.add(BackupRecord(device_id=device.id, device_ip=device.ip, file_path=str(previous_path), status="success"))
        current = BackupRecord(device_id=device.id, device_ip=device.ip, file_path=str(current_path), status="success")
        db.add(current)
        db.commit()
        backup_id = current.id
    finally:
        db.close()

    response = client.get(f"/backup/{backup_id}/diff", headers=headers)

    assert response.status_code == 404


def test_open_backup_directory_rejects_record_paths_outside_the_backup_root(tmp_path, monkeypatch):
    headers = _bootstrap_and_login()
    allowed_root = tmp_path / "allowed-backups"
    allowed_root.mkdir()
    outside_path = tmp_path / "outside-backup.txt"
    outside_path.write_text("external data", encoding="utf-8")
    monkeypatch.setattr(
        backup_router,
        "get_settings",
        lambda _db: SimpleNamespace(backup_directory=str(allowed_root)),
    )
    opened = []
    monkeypatch.setattr(backup_router, "_open_directory", lambda directory: opened.append(directory))
    db = SessionLocal()
    try:
        device = Device(name="Unsafe-Open", ip="10.99.70.12", vendor="Huawei", device_type="switch")
        db.add(device)
        db.flush()
        record = BackupRecord(
            device_id=device.id,
            device_ip=device.ip,
            file_path=str(outside_path),
            status="success",
        )
        db.add(record)
        db.commit()
        backup_id = record.id
    finally:
        db.close()

    response = client.post(f"/backup/{backup_id}/open-directory", headers=headers)

    assert response.status_code == 404
    assert opened == []


def test_backup_diff_ignores_configured_dynamic_lines(tmp_path):
    headers = _bootstrap_and_login()
    patterns = [r"^# Generated at"]
    settings = client.put(
        "/settings",
        headers=headers,
        json={"backup_ignore_patterns": patterns, "backup_directory": str(tmp_path)},
    )

    assert settings.status_code == 200
    assert settings.json().get("backup_ignore_patterns") == patterns
    previous_path = tmp_path / "dynamic-previous.txt"
    current_path = tmp_path / "dynamic-current.txt"
    previous_path.write_text("sysname Core-SW\n# Generated at 2026-08-01 08:00\n", encoding="utf-8")
    current_path.write_text("sysname Core-SW\n# Generated at 2026-08-02 08:00\n", encoding="utf-8")
    db = SessionLocal()
    try:
        device = Device(name="Dynamic-Diff-SW", ip="10.99.70.2", vendor="Huawei", device_type="switch")
        db.add(device)
        db.flush()
        db.add(BackupRecord(device_id=device.id, device_ip=device.ip, file_path=str(previous_path), status="success"))
        current = BackupRecord(device_id=device.id, device_ip=device.ip, file_path=str(current_path), status="success")
        db.add(current)
        db.commit()
        backup_id = current.id
    finally:
        db.close()

    response = client.get(f"/backup/{backup_id}/diff", headers=headers)

    assert response.status_code == 200
    assert response.json()["diff"] == []


def test_if_mib_columns_are_summarized_without_fixed_health_values():
    from app.services.collectors import build_interface_metrics

    metrics = build_interface_metrics({
        "name": {"1": "GigabitEthernet0/0/1", "2": "GigabitEthernet0/0/2"},
        "oper_status": {"1": "1", "2": "2"},
        "in_errors": {"1": "3", "2": "0"},
        "out_errors": {"1": "2", "2": "0"},
        "in_discards": {"1": "1", "2": "0"},
        "out_discards": {"1": "0", "2": "4"},
        "in_octets": {"1": "123456", "2": "789"},
        "out_octets": {"1": "654321", "2": "987"},
    })

    assert metrics["interfaces_up"] == 1
    assert metrics["interfaces_down"] == 1
    assert metrics["interface_error_count"] == 5
    assert metrics["interface_discard_count"] == 5
    assert metrics["interfaces"][0]["in_octets"] == 123456
    assert metrics["interfaces"][1]["oper_status"] == "down"


def test_missing_if_mib_error_columns_do_not_fabricate_zero_errors():
    from app.services.collectors import build_interface_metrics

    metrics = build_interface_metrics({
        "name": {"1": "GigabitEthernet0/0/1"},
        "oper_status": {"1": "1"},
    })

    assert "interface_error_count" not in metrics
    assert "interface_discard_count" not in metrics
    assert metrics["interfaces"][0]["in_errors"] is None


def test_auto_mode_keeps_successful_snmp_result_when_ssh_is_not_configured(monkeypatch):
    device = Device(
        name="SNMP-Only",
        ip="10.99.80.1",
        vendor="Cisco",
        device_type="router",
        snmp_version="v2c",
        snmp_community="encrypted-not-used-by-test",
    )
    task = InspectionTask(name="SNMP优先", collection_mode="auto", inspection_level="quick")

    async def fake_collect_snmp(*_args, **_kwargs):
        return ProtocolResult(ok=True, source="SNMP", metrics={"hostname": "SNMP-Only", "uptime": "1d"})

    async def fake_collect_ssh(*_args, **_kwargs):
        return ProtocolResult(ok=False, source="SSH", failure_stage="credentials", message="SSH credentials not configured")

    monkeypatch.setattr(inspection, "collect_snmp", fake_collect_snmp)
    monkeypatch.setattr(inspection, "collect_ssh", fake_collect_ssh)

    result = asyncio.run(inspection.inspect_device(task, device, AppSettings()))

    assert result.conclusion == "normal"
    assert result.collection_source == "SNMP"
    assert "SSH" in (result.message or "")


def test_auto_mode_reports_failed_snmp_when_ssh_is_not_configured(monkeypatch):
    device = Device(
        name="SNMP-Only-Failed",
        ip="10.99.80.9",
        vendor="Huawei",
        device_type="switch",
        snmp_version="v2c",
        snmp_community="encrypted-not-used-by-test",
    )
    task = InspectionTask(name="SNMP失败", collection_mode="auto", inspection_level="quick")

    async def fake_collect_snmp(*_args, **_kwargs):
        return ProtocolResult(
            ok=False,
            source="SNMP",
            failure_stage="connection",
            message="No SNMP response received before timeout",
        )

    async def unexpected_collect_ssh(*_args, **_kwargs):
        raise AssertionError("SSH collector must not run without configured credentials")

    monkeypatch.setattr(inspection, "collect_snmp", fake_collect_snmp)
    monkeypatch.setattr(inspection, "collect_ssh", unexpected_collect_ssh)

    result = asyncio.run(inspection.inspect_device(task, device, AppSettings()))

    assert result.conclusion == "connection_failed"
    assert result.collection_source == "SNMP"
    assert result.failure_stage == "connection"
    assert result.message == (
        "SNMP failed, SSH not configured: No SNMP response received before timeout"
    )


def test_snmp_interface_collection_failure_is_not_reported_as_normal(monkeypatch):
    device = Device(
        name="SNMP-Interface-Failure",
        ip="10.99.80.2",
        vendor="Cisco",
        device_type="router",
        snmp_version="v2c",
        snmp_community="encrypted-not-used-by-test",
    )
    task = InspectionTask(name="接口采集失败", collection_mode="snmp", inspection_level="quick")

    async def fake_collect_snmp(*_args, **_kwargs):
        return ProtocolResult(
            ok=True,
            source="SNMP",
            metrics={
                "hostname": "SNMP-Interface-Failure",
                "interface_collection_status": "data_fetch_failed",
                "interface_collection_message": "ifTable walk timed out",
            },
        )

    monkeypatch.setattr(inspection, "collect_snmp", fake_collect_snmp)

    result = asyncio.run(inspection.inspect_device(task, device, AppSettings()))

    assert result.conclusion == "data_fetch_failed"
    assert result.collection_status == "failed"
    assert result.failure_stage == "collection"
    assert result.message == "ifTable walk timed out"
    assert result.metrics["hostname"] == "SNMP-Interface-Failure"


def test_ssh_parser_extracts_only_metrics_present_in_vendor_output():
    from app.services.parsers import parse_ssh_metrics

    metrics = parse_ssh_metrics(
        "Huawei",
        """$ display current-configuration | include sysname
sysname Core-SW-01

$ display version
Core-SW-01 uptime is 12 weeks, 3 days

$ display cpu-usage
CPU utilization for five seconds: 73%

$ display memory-usage
Memory Using Percentage Is: 61%

$ display stp brief
0 GigabitEthernet1/0/1 ROOT FORWARDING NONE
0 GigabitEthernet1/0/2 ALTERNATE DISCARDING NONE
""",
    )

    assert metrics["hostname"] == "Core-SW-01"
    assert metrics["uptime"] == "12 weeks, 3 days"
    assert metrics["cpu"] == 73
    assert metrics["memory"] == 61
    assert metrics["stp_discarding_ports"] == 1
    assert parse_ssh_metrics("Cisco", "$ show version\nno recognized metrics\n") == {}


def test_database_backup_is_valid_sqlite_and_can_be_restored():
    headers = _bootstrap_and_login()

    backup = client.get("/system/database-backup", headers=headers)

    assert backup.status_code == 200
    assert backup.content.startswith(b"SQLite format 3\x00")
    restored = client.post(
        "/system/database-restore",
        headers=headers,
        data={"confirmed": "true"},
        files={"file": ("inspection-backup.db", backup.content, "application/x-sqlite3")},
    )
    assert restored.status_code == 200
    assert restored.json()["ok"] is True
    assert client.get("/auth/status").json()["initialized"] is True
    assert client.get("/settings", headers=headers).status_code == 401


def test_database_backup_download_removes_its_temporary_snapshot(tmp_path, monkeypatch):
    headers = _bootstrap_and_login()
    monkeypatch.setattr(system_router, "DATABASE_BACKUP_DIR", tmp_path)

    response = client.get("/system/database-backup", headers=headers)

    assert response.status_code == 200
    assert response.content.startswith(b"SQLite format 3\x00")
    assert list(tmp_path.glob("*.db")) == []


def test_successful_restore_retains_a_readable_pre_restore_snapshot(tmp_path, monkeypatch):
    headers = _bootstrap_and_login()
    monkeypatch.setattr(system_router, "DATABASE_BACKUP_DIR", tmp_path)
    target = client.get("/system/database-backup", headers=headers)
    assert target.status_code == 200
    marker_ip = "198.51.100.252"
    with SessionLocal() as db:
        db.add(Device(ip=marker_ip, name="Pre-restore marker", vendor="Huawei", device_type="switch"))
        db.commit()

    response = client.post(
        "/system/database-restore",
        headers=headers,
        data={"confirmed": "true"},
        files={"file": ("target.db", target.content, "application/vnd.sqlite3")},
    )

    assert response.status_code == 200
    snapshots = list(tmp_path.glob("pre_restore_*.db"))
    assert len(snapshots) == 1
    with closing(sqlite3.connect(snapshots[0])) as connection:
        assert connection.execute("PRAGMA integrity_check").fetchone() == ("ok",)
        assert connection.execute("SELECT ip FROM devices WHERE ip = ?", (marker_ip,)).fetchone() == (marker_ip,)
    with SessionLocal() as db:
        assert db.query(Device).filter_by(ip=marker_ip).first() is None
    assert client.get("/settings", headers=headers).status_code == 401


def test_restore_migrates_previous_version_backup_roots(tmp_path, monkeypatch):
    headers = _bootstrap_and_login()
    old_root = tmp_path / "old-root"
    old_root.mkdir()
    incoming = tmp_path / "previous-version.db"
    system_router.create_database_snapshot(incoming)
    historical_path = old_root / "history.txt"
    historical_path.write_text("sysname Restored-history\n", encoding="utf-8")
    with closing(sqlite3.connect(incoming)) as connection:
        connection.execute(
            "UPDATE app_settings SET backup_directory = ? WHERE id = 1", (str(old_root),),
        )
        connection.execute("DELETE FROM backup_records")
        connection.execute(
            "INSERT INTO backup_records (device_ip, file_path, status, created_at) "
            "VALUES (?, ?, 'success', CURRENT_TIMESTAMP)",
            ("198.51.100.253", str(historical_path)),
        )
        connection.execute("ALTER TABLE backup_records DROP COLUMN backup_root")
        connection.commit()
    snapshot_dir = tmp_path / "snapshots"
    snapshot_dir.mkdir()
    monkeypatch.setattr(system_router, "DATABASE_BACKUP_DIR", snapshot_dir)
    response = client.post(
        "/system/database-restore", headers=headers, data={"confirmed": "true"},
        files={"file": ("previous-version.db", incoming.read_bytes(), "application/vnd.sqlite3")},
    )

    assert response.status_code == 200
    with SessionLocal() as db:
        record = db.query(BackupRecord).filter_by(device_ip="198.51.100.253").one()
        assert record.backup_root == str(old_root.resolve())
        assert record.file_path == str(historical_path.resolve())
    assert len(list(snapshot_dir.glob("pre_restore_*.db"))) == 1


def test_database_restore_rolls_back_when_reinitialization_fails(monkeypatch):
    headers = _bootstrap_and_login()
    target_backup = client.get("/system/database-backup", headers=headers)
    marker = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Restore rollback marker",
            "ip": "10.99.91.249",
            "vendor": "Huawei",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "readonly",
        },
    )
    assert marker.status_code == 201
    real_init_db = system_router.init_db
    calls = 0

    def fail_once_then_initialize():
        nonlocal calls
        calls += 1
        if calls == 1:
            raise RuntimeError("simulated initialization failure")
        real_init_db()

    monkeypatch.setattr(system_router, "init_db", fail_once_then_initialize)

    with pytest.raises(RuntimeError, match="simulated initialization failure"):
        client.post(
            "/system/database-restore",
            headers=headers,
            data={"confirmed": "true"},
            files={"file": ("target.db", target_backup.content, "application/vnd.sqlite3")},
        )

    db = SessionLocal()
    try:
        assert db.query(Device).filter(Device.ip == "10.99.91.249").one_or_none() is not None
    finally:
        db.close()
    assert calls == 2
    assert client.get("/settings", headers=headers).status_code == 200


def test_database_restore_rolls_back_when_restore_audit_commit_fails(tmp_path, monkeypatch):
    headers = _bootstrap_and_login()
    target_backup = client.get("/system/database-backup", headers=headers)
    marker = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Restore audit rollback marker",
            "ip": "10.99.91.248",
            "vendor": "Huawei",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "readonly",
        },
    )
    assert marker.status_code == 201
    real_session_factory = system_router.SessionLocal

    def failing_audit_session_factory():
        audit_db = real_session_factory()

        def fail_commit():
            raise RuntimeError("simulated restore audit commit failure")

        audit_db.commit = fail_commit
        return audit_db

    monkeypatch.setattr(system_router, "DATABASE_BACKUP_DIR", tmp_path)
    monkeypatch.setattr(system_router, "SessionLocal", failing_audit_session_factory)

    with pytest.raises(RuntimeError, match="restore audit commit failure"):
        client.post(
            "/system/database-restore",
            headers=headers,
            data={"confirmed": "true"},
            files={"file": ("target.db", target_backup.content, "application/vnd.sqlite3")},
        )

    db = SessionLocal()
    try:
        assert db.query(Device).filter(Device.ip == "10.99.91.248").one_or_none() is not None
    finally:
        db.close()
    assert list(tmp_path.iterdir()) == []


def test_database_restore_rejects_when_an_inspection_task_is_active():
    headers = _bootstrap_and_login()
    backup = client.get("/system/database-backup", headers=headers)
    db = SessionLocal()
    try:
        db.add(InspectionTask(
            name="Active task blocks restore",
            inspection_level="quick",
            collection_mode="snmp",
            status="running",
            total_devices=1,
        ))
        db.commit()
        task_id = db.query(InspectionTask).order_by(InspectionTask.id.desc()).first().id
    finally:
        db.close()

    response = client.post(
        "/system/database-restore",
        headers=headers,
        data={"confirmed": "true"},
        files={"file": ("inspection-backup.db", backup.content, "application/x-sqlite3")},
    )

    assert response.status_code == 409
    assert response.json()["detail"] == "Cannot restore while an inspection task is running"
    assert client.get("/settings", headers=headers).status_code == 200
    db = SessionLocal()
    try:
        db.get(InspectionTask, task_id).status = "failed"
        db.commit()
    finally:
        db.close()


def test_database_restore_rejects_when_a_configuration_backup_is_active():
    headers = _bootstrap_and_login()
    backup = client.get("/system/database-backup", headers=headers)
    try:
        maintenance = importlib.import_module("app.services.maintenance")
        backup_operation = getattr(maintenance, "backup_operation", None)
    except ModuleNotFoundError:
        backup_operation = None
    assert callable(backup_operation)

    async def restore_while_backup_is_active():
        async with backup_operation():
            return await asyncio.to_thread(
                client.post,
                "/system/database-restore",
                headers=headers,
                data={"confirmed": "true"},
                files={"file": ("inspection-backup.db", backup.content, "application/x-sqlite3")},
            )

    response = asyncio.run(restore_while_backup_is_active())

    assert response.status_code == 409
    assert response.json()["detail"] == "Cannot restore while a configuration backup is running"


def test_task_submission_waits_until_database_restore_validation_finishes(monkeypatch):
    headers = _bootstrap_and_login()
    device = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Restore-race-device",
            "ip": "10.99.91.250",
            "vendor": "Huawei",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "readonly",
        },
    )
    assert device.status_code == 201
    backup = client.get("/system/database-backup", headers=headers)
    restore_entered = Event()
    release_restore = Event()
    task_finished = Event()
    responses = {}

    def block_restore_validation(_path):
        restore_entered.set()
        assert release_restore.wait(5)
        raise HTTPException(status_code=400, detail="Validation stopped by test")

    async def skip_execution(*_args, **_kwargs):
        return None

    monkeypatch.setattr(system_router, "validate_database_backup", block_restore_validation)
    monkeypatch.setattr(tasks_router, "execute_task", skip_execution)

    def restore_request():
        responses["restore"] = client.post(
            "/system/database-restore",
            headers=headers,
            data={"confirmed": "true"},
            files={"file": ("inspection-backup.db", backup.content, "application/x-sqlite3")},
        )

    def task_request():
        responses["task"] = client.post(
            "/tasks",
            headers=headers,
            json={
                "name": "Submitted during restore",
                "device_ids": [device.json()["id"]],
                "inspection_level": "quick",
                "collection_mode": "snmp",
            },
        )
        task_finished.set()

    restore_thread = Thread(target=restore_request)
    restore_thread.start()
    assert restore_entered.wait(5)
    task_thread = Thread(target=task_request)
    task_thread.start()
    try:
        assert not task_finished.wait(0.25), "task creation must wait while database restore owns the maintenance lock"
    finally:
        release_restore.set()
        restore_thread.join(5)
        task_thread.join(5)

    assert responses["restore"].status_code == 400
    assert responses["task"].status_code == 202
    db = SessionLocal()
    try:
        task = db.get(InspectionTask, responses["task"].json()["id"])
        task.status = "failed"
        task.finished_at = datetime.utcnow()
        db.commit()
    finally:
        db.close()


def test_settings_update_cannot_report_success_while_database_restore_replaces_it(monkeypatch):
    headers = _bootstrap_and_login()
    backup = client.get("/system/database-backup", headers=headers)
    replace_entered = Event()
    release_replace = Event()
    settings_finished = Event()
    responses = {}
    real_replace = system_router.replace_database_file

    def blocked_replace(source, destination):
        replace_entered.set()
        assert release_replace.wait(5)
        real_replace(source, destination)

    monkeypatch.setattr(system_router, "replace_database_file", blocked_replace)

    def restore_request():
        responses["restore"] = client.post(
            "/system/database-restore",
            headers=headers,
            data={"confirmed": "true"},
            files={"file": ("inspection-backup.db", backup.content, "application/x-sqlite3")},
        )

    def settings_request():
        responses["settings"] = client.put(
            "/settings", headers=headers, json={"max_concurrency": 17}
        )
        settings_finished.set()

    restore_thread = Thread(target=restore_request)
    settings_thread = Thread(target=settings_request)
    restore_thread.start()
    assert replace_entered.wait(5)
    settings_thread.start()
    try:
        assert not settings_finished.wait(0.25)
    finally:
        release_replace.set()
        restore_thread.join(5)
        settings_thread.join(5)

    assert responses["restore"].status_code == 200
    assert responses["settings"].status_code in {200, 401}
    if responses["settings"].status_code == 200:
        db = SessionLocal()
        try:
            assert db.get(AppSettings, 1).max_concurrency == 17
        finally:
            db.close()


def test_database_restore_requires_explicit_confirmation():
    headers = _bootstrap_and_login()
    backup = client.get("/system/database-backup", headers=headers)

    response = client.post(
        "/system/database-restore",
        headers=headers,
        files={"file": ("inspection-backup.db", backup.content, "application/x-sqlite3")},
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "Database restore requires confirmation"


def test_unconfirmed_restore_audit_waits_for_active_database_restore():
    headers = _bootstrap_and_login()

    async def request_while_restore_is_active():
        async with maintenance.database_restore_operation():
            request = asyncio.create_task(asyncio.to_thread(
                client.post,
                "/system/database-restore",
                headers=headers,
                data={"confirmed": "false"},
                files={"file": ("ignored.db", b"not used", "application/x-sqlite3")},
            ))
            await asyncio.sleep(0.1)
            assert not request.done(), "the unconfirmed restore audit write must register database access"
        return await request

    response = asyncio.run(request_while_restore_is_active())
    assert response.status_code == 400


def test_report_export_saves_timestamped_copy_and_records_log(tmp_path):
    headers = _bootstrap_and_login()
    db = SessionLocal()
    try:
        settings = db.query(AppSettings).first()
        if not settings:
            settings = AppSettings()
            db.add(settings)
        settings.report_directory = str(tmp_path)
        task = InspectionTask(
            name="Report export task",
            inspection_level="quick",
            collection_mode="snmp",
            status="completed",
            total_devices=1,
            completed_devices=1,
            success_devices=1,
            finished_at=datetime.utcnow(),
        )
        task.results.append(InspectionResult(
            device_id=None,
            device_ip="10.99.91.1",
            device_name="Report-SW-01",
            vendor="Huawei",
            collection_source="SNMP",
            connection_status="success",
            authentication_status="success",
            collection_status="success",
            parsing_status="success",
            conclusion="normal",
            failure_stage=None,
            message=None,
            metrics={"uptime": "1 day"},
            raw_output=None,
        ))
        db.add(task)
        db.commit()
        task_id = task.id
    finally:
        db.close()

    response = client.get(f"/excel/report/{task_id}", headers=headers)

    assert response.status_code == 200
    saved_reports = list(tmp_path.glob("Report_export_task_*.xlsx"))
    assert len(saved_reports) == 1
    assert saved_reports[0].read_bytes().startswith(b"PK")
    db = SessionLocal()
    try:
        assert db.query(ActivityLog).filter_by(event="inspection_report_exported", task_id=task_id).count() == 1
    finally:
        db.close()


def test_report_export_rejects_a_task_that_is_not_finished(tmp_path):
    headers = _bootstrap_and_login()
    db = SessionLocal()
    try:
        settings = db.query(AppSettings).first()
        if not settings:
            settings = AppSettings()
            db.add(settings)
        settings.report_directory = str(tmp_path)
        task = InspectionTask(
            name="Running task cannot export",
            inspection_level="quick",
            collection_mode="snmp",
            status="running",
            total_devices=1,
        )
        db.add(task)
        db.commit()
        task_id = task.id
    finally:
        db.close()

    response = client.get(f"/excel/report/{task_id}", headers=headers)

    assert response.status_code == 409
    assert response.json()["detail"] == "Inspection report is only available after task completion"
    assert not list(tmp_path.glob("*.xlsx"))


def test_report_export_rejects_a_cancelled_task_until_the_worker_finishes(tmp_path):
    headers = _bootstrap_and_login()
    db = SessionLocal()
    try:
        settings = db.query(AppSettings).first()
        if not settings:
            settings = AppSettings()
            db.add(settings)
        settings.report_directory = str(tmp_path)
        task = InspectionTask(
            name="Cancelled task still draining",
            inspection_level="quick",
            collection_mode="snmp",
            status="cancelled",
            total_devices=2,
            completed_devices=1,
            finished_at=None,
        )
        db.add(task)
        db.commit()
        task_id = task.id
    finally:
        db.close()

    response = client.get(f"/excel/report/{task_id}", headers=headers)

    assert response.status_code == 409
    assert response.json()["detail"] == "Inspection report is only available after task completion"
    assert not list(tmp_path.glob("*.xlsx"))


def test_database_restore_rejects_while_a_cancelled_task_is_still_finishing():
    headers = _bootstrap_and_login()
    backup = client.get("/system/database-backup", headers=headers)
    db = SessionLocal()
    try:
        task = InspectionTask(
            name="Cancelled task still blocks restore",
            inspection_level="quick",
            collection_mode="snmp",
            status="cancelled",
            total_devices=1,
            finished_at=None,
        )
        db.add(task)
        db.commit()
        task_id = task.id
    finally:
        db.close()

    response = client.post(
        "/system/database-restore",
        headers=headers,
        data={"confirmed": "true"},
        files={"file": ("inspection-backup.db", backup.content, "application/x-sqlite3")},
    )

    assert response.status_code == 409
    assert response.json()["detail"] == "Cannot restore while an inspection task is running"
    db = SessionLocal()
    try:
        task = db.get(InspectionTask, task_id)
        task.finished_at = datetime.utcnow()
        db.commit()
    finally:
        db.close()


def test_report_export_writes_a_temporary_file_before_the_final_xlsx(tmp_path, monkeypatch):
    headers = _bootstrap_and_login()
    db = SessionLocal()
    try:
        settings = db.query(AppSettings).first()
        if not settings:
            settings = AppSettings()
            db.add(settings)
        settings.report_directory = str(tmp_path)
        task = InspectionTask(
            name="Atomic report task",
            inspection_level="quick",
            collection_mode="snmp",
            status="completed",
            total_devices=1,
            completed_devices=1,
            success_devices=1,
            finished_at=datetime.utcnow(),
        )
        task.results.append(InspectionResult(
            device_id=None,
            device_ip="10.99.91.11",
            device_name="Atomic-Report-SW",
            vendor="Huawei",
            collection_source="SNMP",
            connection_status="success",
            authentication_status="success",
            collection_status="success",
            parsing_status="success",
            conclusion="normal",
            failure_stage=None,
            message=None,
            metrics={},
            raw_output=None,
        ))
        db.add(task)
        db.commit()
        task_id = task.id
    finally:
        db.close()

    written_paths: list[str] = []
    original_write_bytes = Path.write_bytes

    def capture_write_path(path, content):
        written_paths.append(path.name)
        return original_write_bytes(path, content)

    monkeypatch.setattr(Path, "write_bytes", capture_write_path)
    response = client.get(f"/excel/report/{task_id}", headers=headers)

    assert response.status_code == 200
    assert any(name.endswith(".tmp") for name in written_paths)
    assert not list(tmp_path.glob("*.tmp"))
    assert (tmp_path / response.headers["content-disposition"].split('filename="', 1)[1].rstrip('"')).is_file()


def test_report_export_rejects_insufficient_disk_space_before_writing(tmp_path, monkeypatch):
    headers = _bootstrap_and_login()
    db = SessionLocal()
    try:
        settings = db.query(AppSettings).first()
        if not settings:
            settings = AppSettings()
            db.add(settings)
        settings.report_directory = str(tmp_path)
        task = InspectionTask(
            name="Disk space report task",
            inspection_level="quick",
            collection_mode="snmp",
            status="completed",
            total_devices=1,
            completed_devices=1,
            success_devices=1,
            finished_at=datetime.utcnow(),
        )
        task.results.append(InspectionResult(
            device_id=None,
            device_ip="10.99.91.10",
            device_name="Disk-Report-SW",
            vendor="Huawei",
            collection_source="SNMP",
            connection_status="success",
            authentication_status="success",
            collection_status="success",
            parsing_status="success",
            conclusion="normal",
            failure_stage=None,
            message=None,
            metrics={},
            raw_output=None,
        ))
        db.add(task)
        db.commit()
        task_id = task.id
    finally:
        db.close()

    def no_space(*_args, **_kwargs):
        raise OSError(errno.ENOSPC, "No space left on device")

    monkeypatch.setattr(excel_router, "ensure_available_space", no_space, raising=False)
    response = client.get(f"/excel/report/{task_id}", headers=headers)

    assert response.status_code == 507
    assert response.json()["detail"] == "磁盘可用空间不足，无法保存报告"
    assert not list(tmp_path.glob("*.xlsx"))
    db = SessionLocal()
    try:
        assert db.query(ActivityLog).filter_by(event="inspection_report_export_failed", task_id=task_id).count() == 1
    finally:
        db.close()


def test_screenshot_upload_saves_png_and_records_log(tmp_path):
    headers = _bootstrap_and_login()
    db = SessionLocal()
    try:
        settings = db.query(AppSettings).first()
        if not settings:
            settings = AppSettings()
            db.add(settings)
        settings.report_directory = str(tmp_path)
        task = InspectionTask(
            name="Screenshot task",
            inspection_level="quick",
            collection_mode="ssh",
            status="completed",
            total_devices=1,
            completed_devices=1,
            success_devices=1,
        )
        task.results.append(InspectionResult(
            device_id=None,
            device_ip="10.99.91.2",
            device_name="Screenshot-SW-01",
            vendor="Cisco",
            collection_source="SSH",
            connection_status="success",
            authentication_status="success",
            collection_status="success",
            parsing_status="success",
            conclusion="normal",
            failure_stage=None,
            message=None,
            metrics={"cpu": 12},
            raw_output=None,
        ))
        db.add(task)
        db.commit()
        task_id = task.id
        result_id = task.results[0].id
    finally:
        db.close()

    png = b"\x89PNG\r\n\x1a\ninspection-result"
    response = client.post(
        "/screenshots",
        headers=headers,
        data={"task_id": str(task_id), "result_id": str(result_id)},
        files={"image": ("result.png", png, "image/png")},
    )

    assert response.status_code == 201
    saved_path = tmp_path / response.json()["filename"]
    assert saved_path.read_bytes() == png
    db = SessionLocal()
    try:
        assert db.query(ActivityLog).filter_by(event="result_screenshot_saved", task_id=task_id).count() == 1
    finally:
        db.close()


def test_screenshot_upload_writes_a_temporary_file_before_the_final_png(tmp_path, monkeypatch):
    headers = _bootstrap_and_login()
    db = SessionLocal()
    try:
        settings = db.query(AppSettings).first()
        if not settings:
            settings = AppSettings()
            db.add(settings)
        settings.report_directory = str(tmp_path)
        task = InspectionTask(
            name="Atomic screenshot task",
            inspection_level="quick",
            collection_mode="snmp",
            status="completed",
            total_devices=1,
            completed_devices=1,
            success_devices=1,
        )
        task.results.append(InspectionResult(
            device_id=None,
            device_ip="10.99.91.20",
            device_name="Atomic-Screenshot-SW",
            vendor="Cisco",
            collection_source="SNMP",
            connection_status="success",
            authentication_status="success",
            collection_status="success",
            parsing_status="success",
            conclusion="normal",
            failure_stage=None,
            message=None,
            metrics={},
            raw_output=None,
        ))
        db.add(task)
        db.commit()
        task_id = task.id
        result_id = task.results[0].id
    finally:
        db.close()

    written_paths: list[str] = []
    original_write_bytes = Path.write_bytes

    def capture_write_path(path, content):
        written_paths.append(path.name)
        return original_write_bytes(path, content)

    monkeypatch.setattr(Path, "write_bytes", capture_write_path)
    response = client.post(
        "/screenshots",
        headers=headers,
        data={"task_id": str(task_id), "result_id": str(result_id)},
        files={"image": ("result.png", b"\x89PNG\r\n\x1a\natomic", "image/png")},
    )

    assert response.status_code == 201
    assert any(name.endswith(".tmp") for name in written_paths)
    assert not list(tmp_path.glob("*.tmp"))
    assert (tmp_path / response.json()["filename"]).is_file()


def test_screenshot_upload_rejects_insufficient_disk_space_before_writing(tmp_path, monkeypatch):
    headers = _bootstrap_and_login()
    db = SessionLocal()
    try:
        settings = db.query(AppSettings).first()
        if not settings:
            settings = AppSettings()
            db.add(settings)
        settings.report_directory = str(tmp_path)
        task = InspectionTask(
            name="Disk space screenshot task",
            inspection_level="quick",
            collection_mode="snmp",
            status="completed",
            total_devices=1,
            completed_devices=1,
            success_devices=1,
        )
        task.results.append(InspectionResult(
            device_id=None,
            device_ip="10.99.91.21",
            device_name="Disk-Screenshot-SW",
            vendor="Cisco",
            collection_source="SNMP",
            connection_status="success",
            authentication_status="success",
            collection_status="success",
            parsing_status="success",
            conclusion="normal",
            failure_stage=None,
            message=None,
            metrics={},
            raw_output=None,
        ))
        db.add(task)
        db.commit()
        task_id = task.id
        result_id = task.results[0].id
    finally:
        db.close()

    def no_space(*_args, **_kwargs):
        raise OSError(errno.ENOSPC, "No space left on device")

    monkeypatch.setattr(screenshots_router, "ensure_available_space", no_space, raising=False)
    response = client.post(
        "/screenshots",
        headers=headers,
        data={"task_id": str(task_id), "result_id": str(result_id)},
        files={"image": ("result.png", b"\x89PNG\r\n\x1a\ndisk-space", "image/png")},
    )

    assert response.status_code == 507
    assert response.json()["detail"] == "磁盘可用空间不足，无法保存截图"
    assert not list(tmp_path.glob("*.png"))
    db = SessionLocal()
    try:
        assert db.query(ActivityLog).filter_by(event="result_screenshot_failed", task_id=task_id).count() == 1
    finally:
        db.close()


def test_backup_can_open_its_saved_directory(monkeypatch, tmp_path):
    headers = _bootstrap_and_login()
    client.put("/settings", headers=headers, json={"backup_directory": str(tmp_path)})
    path = tmp_path / "backup.txt"
    path.write_text("display current-configuration", encoding="utf-8")
    db = SessionLocal()
    try:
        device = Device(name="Backup-SW-01", ip="10.99.91.3", vendor="H3C", device_type="switch")
        db.add(device)
        db.flush()
        record = BackupRecord(device_id=device.id, device_ip=device.ip, file_path=str(path), status="success")
        db.add(record)
        db.commit()
        backup_id = record.id
    finally:
        db.close()
    opened: list[str] = []
    monkeypatch.setattr(backup_router, "_open_directory", lambda directory: opened.append(str(directory)), raising=False)
    response = client.post(f"/backup/{backup_id}/open-directory", headers=headers)

    assert response.status_code == 200
    assert opened == [str(tmp_path)]
    db = SessionLocal()
    try:
        assert db.query(ActivityLog).filter_by(event="backup_directory_opened", device_ip="10.99.91.3").count() == 1
    finally:
        db.close()


def test_backup_history_delete_endpoint_removes_file_and_record(tmp_path):
    headers = _bootstrap_and_login()
    client.put("/settings", headers=headers, json={"backup_directory": str(tmp_path)})
    path = tmp_path / "delete-through-api.txt"
    path.write_text("display current-configuration", encoding="utf-8")
    db = SessionLocal()
    try:
        record = BackupRecord(
            device_id=None,
            device_ip="10.99.91.4",
            file_path=str(path),
            status="success",
        )
        db.add(record)
        db.commit()
        backup_id = record.id
    finally:
        db.close()

    response = client.delete(f"/backup/{backup_id}", headers=headers)

    assert response.status_code == 200
    assert response.json() == {"deleted": 1}
    assert not path.exists()
    with SessionLocal() as db:
        assert db.get(BackupRecord, backup_id) is None
        assert db.query(ActivityLog).filter_by(
            event="backup_deleted", device_ip="10.99.91.4"
        ).count() == 1


def test_device_excel_export_excludes_all_credentials():
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Export-Test",
            "ip": "10.99.90.1",
            "vendor": "Ruijie",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "must-not-appear-in-export",
        },
    )
    assert created.status_code == 201

    response = client.get("/excel/devices-export", headers=headers)

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook["设备清单"]
    headers_row = [cell.value for cell in sheet[1]]
    assert "ssh_password" not in headers_row
    assert "snmp_community" not in headers_row
    assert "snmp_v3_auth_key" not in headers_row
    assert any(row[1].value == "10.99.90.1" for row in sheet.iter_rows(min_row=2))


def test_snmp_connection_test_reports_success_without_creating_inspection(monkeypatch):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Connectivity-SNMP",
            "ip": "10.99.90.2",
            "vendor": "Huawei",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "test-community",
        },
    )
    assert created.status_code == 201

    async def fake_snmp_connectivity(*_args, **_kwargs):
        return ProtocolResult(True, "SNMP")

    monkeypatch.setattr(devices_router, "test_snmp_connectivity", fake_snmp_connectivity, raising=False)
    db = SessionLocal()
    try:
        task_count_before = db.query(InspectionTask).count()
    finally:
        db.close()
    response = client.post(
        f"/devices/{created.json()['id']}/connection-test",
        headers=headers,
        json={"protocol": "snmp"},
    )

    assert response.status_code == 200
    assert response.json() == {
        "protocol": "SNMP",
        "status": "success",
        "failure_stage": None,
        "message": None,
    }
    db = SessionLocal()
    try:
        assert db.query(InspectionTask).count() == task_count_before
        log = db.query(ActivityLog).filter_by(event="device_connectivity_test", device_ip="10.99.90.2").one()
        assert log.level == "INFO"
        assert "SNMP success" in (log.message or "")
    finally:
        db.close()


def test_connection_test_returns_authentication_failed_status(monkeypatch):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Connectivity-SSH",
            "ip": "10.99.90.3",
            "vendor": "Cisco",
            "device_type": "router",
            "ssh_username": "ops",
            "ssh_password": "test-password",
        },
    )
    assert created.status_code == 201

    async def fake_ssh_connectivity(*_args, **_kwargs):
        return ProtocolResult(
            False,
            "SSH",
            failure_stage="authentication",
            message="Authentication failed; password=test-password",
            authentication_failed=True,
        )

    monkeypatch.setattr(devices_router, "test_ssh_connectivity", fake_ssh_connectivity, raising=False)
    response = client.post(
        f"/devices/{created.json()['id']}/connection-test",
        headers=headers,
        json={"protocol": "ssh"},
    )

    assert response.status_code == 200
    assert response.json() == {
        "protocol": "SSH",
        "status": "authentication_failed",
        "failure_stage": "authentication",
        "message": "Authentication failed; password [REDACTED]",
    }


def test_task_persists_non_secret_device_identity_snapshot(monkeypatch):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Snapshot-Edge-01",
            "ip": "10.99.90.4",
            "vendor": "H3C",
            "platform_id": "h3c-comware",
            "device_type": "switch",
            "group_name": "Edge",
            "ssh_username": "ops",
            "ssh_password": "snapshot-secret",
        },
    )
    assert created.status_code == 201

    async def skip_execution(*_args, **_kwargs):
        return None

    monkeypatch.setattr(tasks_router, "execute_task", skip_execution)
    response = client.post(
        "/tasks",
        headers=headers,
        json={
            "name": "Snapshot task",
            "device_ids": [created.json()["id"]],
            "inspection_level": "full",
            "collection_mode": "ssh",
        },
    )
    assert response.status_code == 202

    db = SessionLocal()
    try:
        task = db.get(InspectionTask, response.json()["id"])
        assert task.selected_device_snapshot == [{
            "id": created.json()["id"],
            "name": "Snapshot-Edge-01",
            "ip": "10.99.90.4",
            "vendor": "H3C",
            "platform_id": "h3c-comware",
            "device_type": "switch",
            "group_name": "Edge",
        }]
        assert "snapshot-secret" not in str(task.selected_device_snapshot)
    finally:
        db.close()


def test_finished_task_can_be_retried_with_current_settings(monkeypatch):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Retry-Task-SW-01",
            "ip": "10.99.91.41",
            "vendor": "Huawei",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "readonly",
        },
    )

    async def skip_execution(*_args, **_kwargs):
        return None

    monkeypatch.setattr(tasks_router, "execute_task", skip_execution)
    original = client.post(
        "/tasks",
        headers=headers,
        json={
            "name": "Original retry task",
            "device_ids": [created.json()["id"]],
            "inspection_level": "full",
            "collection_mode": "snmp",
        },
    )
    db = SessionLocal()
    try:
        source = db.get(InspectionTask, original.json()["id"])
        source.status = "failed"
        source.finished_at = datetime.utcnow()
        settings = db.query(AppSettings).first()
        settings.max_concurrency = 3
        db.commit()
    finally:
        db.close()

    retried = client.post(f"/tasks/{original.json()['id']}/retry", headers=headers)

    assert retried.status_code == 202
    assert retried.json()["id"] != original.json()["id"]
    assert retried.json()["name"] == "Original retry task（重试）"
    assert retried.json()["inspection_level"] == "full"
    assert retried.json()["collection_mode"] == "snmp"
    db = SessionLocal()
    try:
        task = db.get(InspectionTask, retried.json()["id"])
        assert task.selected_device_ids == str(created.json()["id"])
        assert task.settings_snapshot["max_concurrency"] == 3
    finally:
        db.close()


def test_active_task_cannot_be_retried(monkeypatch):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Active-Retry-SW-01",
            "ip": "10.99.91.42",
            "vendor": "Huawei",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "readonly",
        },
    )

    async def skip_execution(*_args, **_kwargs):
        return None

    monkeypatch.setattr(tasks_router, "execute_task", skip_execution)
    original = client.post(
        "/tasks",
        headers=headers,
        json={"name": "Still active", "device_ids": [created.json()["id"]]},
    )

    response = client.post(f"/tasks/{original.json()['id']}/retry", headers=headers)

    assert response.status_code == 409
    assert response.json()["detail"] == "Only finished tasks can be retried"


def test_finished_task_without_active_original_devices_cannot_be_retried(monkeypatch):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Inactive-Retry-SW-01",
            "ip": "10.99.91.43",
            "vendor": "Huawei",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "readonly",
        },
    )

    async def skip_execution(*_args, **_kwargs):
        return None

    monkeypatch.setattr(tasks_router, "execute_task", skip_execution)
    original = client.post(
        "/tasks",
        headers=headers,
        json={"name": "No active devices", "device_ids": [created.json()["id"]]},
    )
    db = SessionLocal()
    try:
        source = db.get(InspectionTask, original.json()["id"])
        source.status = "completed"
        source.finished_at = datetime.utcnow()
        db.get(Device, created.json()["id"]).is_active = False
        db.commit()
    finally:
        db.close()

    response = client.post(f"/tasks/{original.json()['id']}/retry", headers=headers)

    assert response.status_code == 409
    assert response.json()["detail"] == "No active devices from the original task remain"


def test_task_persists_execution_settings_snapshot(monkeypatch):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Settings-Snapshot-01",
            "ip": "10.99.90.40",
            "vendor": "Huawei",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "readonly",
        },
    )
    assert created.status_code == 201
    updated = client.put(
        "/settings",
        headers=headers,
        json={
            "max_concurrency": 7,
            "snmp_timeout_seconds": 8,
            "snmp_retries": 3,
            "connection_timeout_seconds": 11,
            "command_timeout_seconds": 19,
            "network_retry_count": 4,
            "cpu_warning_threshold": 72,
            "cpu_critical_threshold": 92,
            "memory_warning_threshold": 73,
            "memory_critical_threshold": 93,
            "storage_warning_threshold": 74,
            "storage_critical_threshold": 94,
        },
    )
    assert updated.status_code == 200

    async def skip_execution(*_args, **_kwargs):
        return None

    monkeypatch.setattr(tasks_router, "execute_task", skip_execution)
    response = client.post(
        "/tasks",
        headers=headers,
        json={
            "name": "Settings snapshot task",
            "device_ids": [created.json()["id"]],
            "inspection_level": "quick",
            "collection_mode": "snmp",
        },
    )
    assert response.status_code == 202

    db = SessionLocal()
    try:
        task = db.get(InspectionTask, response.json()["id"])
        assert task.settings_snapshot == {
            "max_concurrency": 7,
            "snmp_timeout_seconds": 8,
            "snmp_retries": 3,
            "connection_timeout_seconds": 11,
            "command_timeout_seconds": 19,
            "network_retry_count": 4,
            "cpu_warning_threshold": 72,
            "cpu_critical_threshold": 92,
            "memory_warning_threshold": 73,
            "memory_critical_threshold": 93,
            "storage_warning_threshold": 74,
            "storage_critical_threshold": 94,
        }
    finally:
        db.close()


def test_task_execution_uses_saved_settings_snapshot_after_settings_change(monkeypatch):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Settings-Execution-01",
            "ip": "10.99.90.41",
            "vendor": "Cisco",
            "device_type": "router",
            "ssh_username": "ops",
            "ssh_password": "not-part-of-the-snapshot",
        },
    )
    assert created.status_code == 201
    initial_settings = {
        "max_concurrency": 6,
        "snmp_timeout_seconds": 7,
        "snmp_retries": 2,
        "connection_timeout_seconds": 12,
        "command_timeout_seconds": 18,
        "network_retry_count": 3,
        "cpu_warning_threshold": 74,
        "memory_warning_threshold": 75,
    }
    assert client.put("/settings", headers=headers, json=initial_settings).status_code == 200

    async def skip_execution(*_args, **_kwargs):
        return None

    monkeypatch.setattr(tasks_router, "execute_task", skip_execution)
    created_task = client.post(
        "/tasks",
        headers=headers,
        json={
            "name": "Settings execution task",
            "device_ids": [created.json()["id"]],
            "inspection_level": "full",
            "collection_mode": "ssh",
        },
    )
    assert created_task.status_code == 202

    changed_settings = {
        "max_concurrency": 1,
        "snmp_timeout_seconds": 1,
        "snmp_retries": 0,
        "connection_timeout_seconds": 2,
        "command_timeout_seconds": 3,
        "network_retry_count": 0,
        "cpu_warning_threshold": 99,
        "cpu_critical_threshold": 100,
        "memory_warning_threshold": 98,
        "memory_critical_threshold": 100,
    }
    assert client.put("/settings", headers=headers, json=changed_settings).status_code == 200
    observed: dict[str, int] = {}

    async def fake_inspect(task, current_device, task_settings):
        observed.update({
            "max_concurrency": task_settings.max_concurrency,
            "snmp_timeout_seconds": task_settings.snmp_timeout_seconds,
            "snmp_retries": task_settings.snmp_retries,
            "connection_timeout_seconds": task_settings.connection_timeout_seconds,
            "command_timeout_seconds": task_settings.command_timeout_seconds,
            "network_retry_count": task_settings.network_retry_count,
            "cpu_warning_threshold": task_settings.cpu_warning_threshold,
            "memory_warning_threshold": task_settings.memory_warning_threshold,
        })
        return InspectionResult(
            task=task,
            device_id=current_device.id,
            device_ip=current_device.ip,
            device_name=current_device.name,
            vendor=current_device.vendor,
            collection_source="SSH",
            connection_status="success",
            authentication_status="success",
            collection_status="success",
            parsing_status="success",
            conclusion="normal",
            failure_stage=None,
            message=None,
            metrics={},
            raw_output=None,
        )

    monkeypatch.setattr(inspection, "inspect_device", fake_inspect)
    asyncio.run(inspection.execute_task(created_task.json()["id"], SessionLocal))

    assert observed == initial_settings


def test_task_execution_uses_captured_encrypted_credentials_after_device_edit(monkeypatch):
    from app.services.security import decrypt_secret, encrypt_secret

    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Credential-Snapshot-01",
            "ip": "10.99.90.43",
            "vendor": "Huawei",
            "device_type": "switch",
            "ssh_username": "ops",
            "ssh_password": "original-readonly-password",
        },
    )
    assert created.status_code == 201

    async def skip_execution(*_args, **_kwargs):
        return None

    monkeypatch.setattr(tasks_router, "execute_task", skip_execution)
    created_task = client.post(
        "/tasks",
        headers=headers,
        json={
            "name": "Credential snapshot task",
            "device_ids": [created.json()["id"]],
            "inspection_level": "quick",
            "collection_mode": "ssh",
        },
    )
    assert created_task.status_code == 202

    db = SessionLocal()
    try:
        device = db.get(Device, created.json()["id"])
        device.ssh_password = encrypt_secret("updated-readonly-password")
        db.commit()
    finally:
        db.close()

    observed: dict[str, str | None] = {}

    async def fake_inspect(task, current_device, task_settings):
        observed["ssh_password"] = decrypt_secret(current_device.ssh_password)
        return InspectionResult(
            task=task,
            device_id=current_device.id,
            device_ip=current_device.ip,
            device_name=current_device.name,
            vendor=current_device.vendor,
            collection_source="SSH",
            connection_status="success",
            authentication_status="success",
            collection_status="success",
            parsing_status="success",
            conclusion="normal",
            failure_stage=None,
            message=None,
            metrics={},
            raw_output=None,
        )

    monkeypatch.setattr(inspection, "inspect_device", fake_inspect)
    asyncio.run(inspection.execute_task(created_task.json()["id"], SessionLocal))

    assert observed["ssh_password"] == "original-readonly-password"


def test_task_execution_keeps_original_device_identity_after_device_update(monkeypatch):
    db = SessionLocal()
    try:
        device = Device(
            name="Original-Edge-01",
            ip="10.99.90.5",
            vendor="Huawei",
            device_type="switch",
            group_name="Edge",
        )
        db.add(device)
        db.flush()
        snapshot = {
            "id": device.id,
            "name": device.name,
            "ip": device.ip,
            "vendor": device.vendor,
            "device_type": device.device_type,
            "group_name": device.group_name,
        }
        task = InspectionTask(
            name="Original identity task",
            inspection_level="quick",
            collection_mode="ssh",
            status="queued",
            total_devices=1,
            selected_device_ids=str(device.id),
            selected_device_snapshot=[snapshot],
        )
        db.add(task)
        db.commit()
        task_id = task.id
        device.name = "Renamed-Edge-01"
        device.ip = "10.99.90.55"
        device.vendor = "Cisco"
        db.commit()
    finally:
        db.close()

    async def fake_inspect(task, current_device, _settings):
        return InspectionResult(
            task=task,
            device_id=current_device.id,
            device_ip=current_device.ip,
            device_name=current_device.name,
            vendor=current_device.vendor,
            collection_source="SSH",
            connection_status="success",
            authentication_status="success",
            collection_status="success",
            parsing_status="success",
            conclusion="normal",
            failure_stage=None,
            message=None,
            metrics={},
            raw_output=None,
        )

    monkeypatch.setattr(inspection, "inspect_device", fake_inspect)
    asyncio.run(inspection.execute_task(task_id, SessionLocal, max_concurrency=1))

    db = SessionLocal()
    try:
        result = db.query(InspectionResult).filter_by(task_id=task_id).one()
        assert result.device_id == snapshot["id"]
        assert result.device_name == "Original-Edge-01"
        assert result.device_ip == "10.99.90.5"
        assert result.vendor == "Huawei"
    finally:
        db.close()


def test_device_update_can_clear_snmp_while_keeping_ssh():
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Dual-protocol-device",
            "ip": "10.99.90.6",
            "vendor": "Cisco",
            "device_type": "router",
            "snmp_version": "v2c",
            "snmp_community": "dual-protocol-community",
            "ssh_username": "ops",
            "ssh_password": "dual-protocol-password",
        },
    )
    assert created.status_code == 201

    response = client.put(
        f"/devices/{created.json()['id']}",
        headers=headers,
        json={"clear_protocols": ["snmp"]},
    )

    assert response.status_code == 200
    assert response.json()["has_snmp_credentials"] is False
    assert response.json()["snmp_version"] is None
    assert response.json()["has_ssh"] is True
    db = SessionLocal()
    try:
        device = db.get(Device, created.json()["id"])
        assert device.snmp_community is None
        assert device.ssh_password is not None
    finally:
        db.close()


def test_changing_ssh_port_clears_the_trusted_host_fingerprint():
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "SSH-port-change",
            "ip": "10.99.90.61",
            "vendor": "Cisco",
            "device_type": "router",
            "ssh_username": "ops",
            "ssh_password": "port-change-password",
        },
    )
    assert created.status_code == 201
    device_id = created.json()["id"]

    db = SessionLocal()
    try:
        device = db.get(Device, device_id)
        device.ssh_host_fingerprint = "SHA256:trusted-on-port-22"
        db.commit()
    finally:
        db.close()

    response = client.put(
        f"/devices/{device_id}",
        headers=headers,
        json={"ssh_port": 2222},
    )

    assert response.status_code == 200
    db = SessionLocal()
    try:
        assert db.get(Device, device_id).ssh_host_fingerprint is None
    finally:
        db.close()


def test_device_update_rejects_removing_final_protocol():
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "SNMP-only-device",
            "ip": "10.99.90.7",
            "vendor": "Ruijie",
            "device_type": "switch",
            "snmp_version": "v2c",
            "snmp_community": "single-protocol-community",
        },
    )
    assert created.status_code == 201

    response = client.put(
        f"/devices/{created.json()['id']}",
        headers=headers,
        json={"clear_protocols": ["snmp"]},
    )

    assert response.status_code == 422
    assert response.json()["detail"] == "At least one complete SNMP or SSH credential set is required"


def test_task_uses_saved_default_collection_mode_when_omitted(monkeypatch):
    headers = _bootstrap_and_login()
    settings = client.put(
        "/settings",
        headers=headers,
        json={"default_collection_mode": "ssh"},
    )
    assert settings.status_code == 200
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Default-mode-device",
            "ip": "10.99.90.8",
            "vendor": "Cisco",
            "device_type": "router",
            "ssh_username": "ops",
            "ssh_password": "default-mode-password",
        },
    )
    assert created.status_code == 201

    async def skip_execution(*_args, **_kwargs):
        return None

    monkeypatch.setattr(tasks_router, "execute_task", skip_execution)
    response = client.post(
        "/tasks",
        headers=headers,
        json={"name": "Uses saved default", "device_ids": [created.json()["id"]]},
    )

    assert response.status_code == 202
    assert response.json()["collection_mode"] == "ssh"


def test_device_import_template_matches_documented_columns():
    """The template must expose every documented column, in order.

    The headers are the Chinese labels the operator reads, so the assertion
    goes through the shared definition instead of repeating one spelling.
    """
    from app.services.device_import_columns import TEMPLATE_COLUMNS, TEMPLATE_HEADERS

    headers = _bootstrap_and_login()

    response = client.get("/excel/template", headers=headers)

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    headers_row = [cell.value for cell in workbook["设备导入模板"][1]]
    assert headers_row == TEMPLATE_HEADERS
    assert [field for field, _label in TEMPLATE_COLUMNS] == [
        "name",
        "ip",
        "vendor",
        "platform_id",
        "device_type",
        "group_name",
        "note",
        "ssh_port",
        "ssh_username",
        "ssh_password",
        "enable_password",
        "snmp_version",
        "snmp_port",
        "snmp_community",
        "snmp_v3_username",
        "snmp_v3_security_level",
        "snmp_v3_auth_protocol",
        "snmp_v3_auth_key",
        "snmp_v3_priv_protocol",
        "snmp_v3_priv_key",
    ]


def test_snmp_retry_budget_merges_protocol_and_network_settings_in_one_collector_call(monkeypatch):
    device = Device(
        name="Retry-SW-01",
        ip="10.99.91.4",
        vendor="Huawei",
        device_type="switch",
    )
    task = InspectionTask(name="Retry test", inspection_level="quick", collection_mode="snmp")
    calls = []

    async def flaky_snmp(*_args, **kwargs):
        calls.append(kwargs)
        return ProtocolResult(False, "SNMP", failure_stage="connection", message="timeout")

    monkeypatch.setattr(inspection, "collect_snmp", flaky_snmp)

    result = asyncio.run(inspection.inspect_device(
        task,
        device,
        AppSettings(network_retry_count=2, snmp_timeout_seconds=3, snmp_retries=4),
    ))

    assert calls == [{"timeout": 3, "retries": 4}]
    assert result.conclusion == "connection_failed"


def test_network_retry_setting_increases_snmp_timeout_retry_budget(monkeypatch):
    device = Device(
        name="Retry-SW-02",
        ip="10.99.91.5",
        vendor="Huawei",
        device_type="switch",
    )
    task = InspectionTask(name="Retry test", inspection_level="quick", collection_mode="snmp")
    calls = []

    async def failed_snmp(*_args, **kwargs):
        calls.append(kwargs)
        return ProtocolResult(False, "SNMP", failure_stage="connection", message="timeout")

    monkeypatch.setattr(inspection, "collect_snmp", failed_snmp)

    result = asyncio.run(inspection.inspect_device(
        task,
        device,
        AppSettings(network_retry_count=2, snmp_timeout_seconds=3, snmp_retries=0),
    ))

    assert calls == [{"timeout": 3, "retries": 2}]
    assert result.conclusion == "connection_failed"


def test_network_retry_uses_default_when_legacy_settings_value_is_missing(monkeypatch):
    device = Device(
        name="Retry-default-SW-01",
        ip="10.99.91.6",
        vendor="Huawei",
        device_type="switch",
        ssh_username="ops",
        ssh_password="encrypted-password",
    )
    task = InspectionTask(name="Retry default", inspection_level="quick", collection_mode="ssh")
    attempts = 0

    async def flaky_ssh(*_args, **_kwargs):
        nonlocal attempts
        attempts += 1
        if attempts < 3:
            return ProtocolResult(False, "SSH", failure_stage="connection", message="timeout")
        return ProtocolResult(True, "SSH", metrics={
            "hostname": "Retry-default-SW-01",
            "model": "S5735",
            "uptime": "1d",
            "cpu": 20,
            "memory": 30,
            "interfaces_total": 24,
            "interface_collection_status": "success",
            "stp_discarding_ports": 0,
            "active_alarms": {"total": 0},
        })

    monkeypatch.setattr(inspection, "collect_ssh", flaky_ssh)

    result = asyncio.run(inspection.inspect_device(
        task,
        device,
        AppSettings(network_retry_count=None),
    ))

    assert attempts == 3
    assert result.conclusion == "normal"


def test_backup_uses_configured_connection_and_command_timeouts(monkeypatch, tmp_path):
    headers = _bootstrap_and_login()
    created = client.post(
        "/devices",
        headers=headers,
        json={
            "name": "Backup-timeout-SW-01",
            "ip": "10.99.91.5",
            "vendor": "Cisco",
            "device_type": "switch",
            "ssh_username": "ops",
            "ssh_password": "backup-timeout-password",
        },
    )
    assert created.status_code == 201
    settings = client.put(
        "/settings",
        headers=headers,
        json={
            "connection_timeout_seconds": 7,
            "command_timeout_seconds": 21,
            "backup_directory": str(tmp_path),
        },
    )
    assert settings.status_code == 200
    observed: list[tuple[int, int]] = []

    async def fake_backup(_device, connection_timeout, command_timeout):
        observed.append((connection_timeout, command_timeout))
        return ProtocolResult(True, "SSH", raw_output="sysname Core-SW\ninterface Vlanif1\nreturn\n")

    monkeypatch.setattr(backup_router, "backup_configuration", fake_backup)
    response = client.post("/backup", headers=headers, json=[created.json()["id"]])

    assert response.status_code == 200
    assert observed == [(7, 21)]
